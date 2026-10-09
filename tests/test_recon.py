"""Tests for the Full Reconciliation engine.

All data here is made up (fake tickets, parts and amounts) so nothing from
a real dealer ever lives in the repo.

Run:  python -m unittest discover -s tests -v
"""
import datetime
import os
import shutil
import sys
import tempfile
import time
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(
    __file__))))

from openpyxl import load_workbook  # noqa: E402

from recon import engine, master as M  # noqa: E402
from recon.common import ocr_ticket  # noqa: E402
from recon.dealerconnect import read_dc_text  # noqa: E402
from recon.memo import parse_memo_text  # noqa: E402
from recon.shipper import (_shipment_from_header, ocr_part,  # noqa: E402
                           parse_line)

MEMO_HEADER = """FCA US LLC
               MOPAR WEEKLY GLOBAL CORE RETURN CREDIT MEMO
                                              BILL TO: 12345000
CREDIT MEMO NUMBER: 03181000CC00900001        ANYTOWN
CREDIT MEMO DATE : OCTOBER 03, 2026           TX 70000
PAGE              :   {page}                       SD NO: 01
LINE PART     DESCRIPTN QTY UNIT         GROSS     EXT   D   DISCNT     NET    P R
"""

MEMO_TEXT = MEMO_HEADER.format(page=1) + """
0318100-1000001 ORD#: GCRS    O/T:M
  1 U1111111AB TRANS       1 2500.00- 2500.00- 99     .0       .00    2500.00- N 1
     REFERENCE/CONTROL NUMBER C400000001
     GCRS PART RETURN CREDIT RM:C400000001 CORE:U1111111AB
                           SUB-TOTAL     2500.00-                     2500.00-

0318100-1000002 ORD#: GCRS    O/T:M
001 ARC31101   DEPOSIT PART VALUES       150.00-                       150.00-
     REFERENCE/CONTROL NUMBER C400000002
     GCRS PART RETURN CREDIT RM:C400000002 CORE:22222222AI
     CLAIM:123456 VIN:1ABCDEFGH12345678 REPAIR:09/09/2026
""" + MEMO_HEADER.format(page=2) + """
                           SUB-TOTAL     150.00-                       150.00-

0318100-1000003 ORD#: GCRS    O/T:M
001 ARC31101   DEPOSIT PART VALUES        75.00-                        75.00-
     REFERENCE/CONTROL NUMBER C400000099
     GCRS PART RETURN CREDIT RM:C400000099 CORE:33333333AA
                           SUB-TOTAL      75.00-                        75.00-
SUMMARY:
    TOTAL GROSS AMOUNT                             2500.00-
    ARC31101 DEPOSIT PART VALUES                    225.00-
    NET CREDIT AMOUNT                              2725.00-
"""


def W(*tokens):
    """Fake OCR words for parse_line."""
    return [{'text': t, 'left': 100 + i * 60, 'top': 400, 'width': 50,
             'height': 20, 'conf': 90} for i, t in enumerate(tokens)]


class TestMemoParser(unittest.TestCase):
    def test_blocks_amounts_and_tie_out(self):
        r = parse_memo_text(MEMO_TEXT, source='t')
        self.assertEqual(r['dealer'], '12345')
        memo = r['memos']['CC00900001']
        self.assertEqual(memo['date'], datetime.date(2026, 10, 3))
        self.assertEqual(memo['net_total'], 2725.0)
        self.assertAlmostEqual(memo['lines_total'], 2725.0)
        self.assertEqual(r['warnings'], [])
        by = {c['ticket']: c for c in r['credits']}
        self.assertEqual(by['C400000001']['amount'], 2500.0)  # > $999
        self.assertEqual(by['C400000001']['part'], 'U1111111AB')
        # sub-total on the next page still belongs to this block
        self.assertEqual(by['C400000002']['amount'], 150.0)
        self.assertEqual(by['C400000002']['claim'], 123456)
        self.assertEqual(by['C400000002']['part'], '22222222AI')  # AI kept
        self.assertEqual(by['C400000099']['amount'], 75.0)

    def test_mismatch_is_reported(self):
        bad = MEMO_TEXT.replace('NET CREDIT AMOUNT                              2725.00-',
                                'NET CREDIT AMOUNT                              2800.00-')
        r = parse_memo_text(bad)
        self.assertTrue(any('NET CREDIT AMOUNT' in w for w in r['warnings']))


class TestOcrHelpers(unittest.TestCase):
    def test_ticket_lookalikes(self):
        self.assertEqual(ocr_ticket('0411100011')[0], 'C411100011')
        self.assertEqual(ocr_ticket('©411100022')[0], 'C411100022')
        self.assertEqual(ocr_ticket('6411100033')[0], 'C411100033')
        self.assertEqual(ocr_ticket('CBO0111108')[0], 'CB00111108')
        self.assertIsNone(ocr_ticket('68520294AB'))     # part, not ticket
        self.assertIsNone(ocr_ticket('04801313AE'))

    def test_part_cleanup(self):
        self.assertEqual(ocr_part('U8i56209AH'), 'U8156209AH')
        self.assertEqual(ocr_part('U8090720AI'), 'U8090720AI')
        self.assertEqual(ocr_part('O5151705AF'), '05151705AF')
        self.assertEqual(ocr_part('$3022263AF'), '53022263AF')
        self.assertEqual(ocr_part('vU8052759AA'), 'U8052759AA')
        self.assertEqual(ocr_part('CEZGU902AB'), 'CEZGU902AB')

    def test_shipper_lines(self):
        r = parse_line(W('411100044', '68474094AA', 'SHAFT', '1-75.00'))
        self.assertEqual((r['ticket'], r['qty'], r['amount']),
                         ('C411100044', 1, 75.0))
        r = parse_line(W('C411100055', 'U8444855AB', 'CORE', '1125.00'))
        self.assertEqual((r['qty'], r['amount']), (1, 125.0))
        r = parse_line(W('UnPaid', 'C411100066', '123987', 'CEZGU902AB',
                         'MANIFOLD', '1', '500.00'))
        self.assertEqual((r['ticket'], r['claim'], r['part'], r['amount']),
                         ('C411100066', 123987, 'CEZGU902AB', 500.0))
        r = parse_line(W('C411100077', 'som', '57009571AA', 'COMPRESSOR',
                         '1', '50.00', ';'))
        self.assertEqual((r['part'], r['amount']), ('57009571AA', 50.0))
        r = parse_line(W('C411100088', 'U8429324AA', 'CORE', '1',
                         '2,000.00'))
        self.assertEqual(r['amount'], 2000.0)
        self.assertIsNone(parse_line(W('PART', 'CONTROL', 'CLAIM')))

    def test_shp_header(self):
        f = _shipment_from_header
        self.assertEqual(f('GCRS DEALER SHIPPER DOCUMENT GCRS SHP#2026-09-25 '
                           '16:30.01'), '2026-09-25 16:30.01')
        self.assertEqual(f('GCRS SHP#42026-08-2I1 18:00.01'),
                         '2026-08-21 18:00.01')
        self.assertEqual(f('GCRS SHP#2026-~-09-21 17:00.01'),
                         '2026-09-21 17:00.01')
        self.assertEqual(f('GCRS SHP#2026-08-03'), '2026-08-03')
        self.assertIsNone(f('GCRS SHP#2026-13-03'))


def make_master(path, rows):
    M.write_master(rows, path)


def row(memo, shp, ticket, part, amount, claim=None, req=None):
    return {'memo': memo, 'shipment': shp, 'ticket': ticket, 'part': part,
            'claim': claim, 'amount': amount, 'requested': req}


class TestEngine(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.asof = datetime.date(2026, 10, 5)
        self.rows = [
            row('CC00800000', None, 'C300000001', '11111111AA', 50),
            row('#000900001', '2026-07-01 17:00.01', 'C400000099',
                '33333333AA', 75),                     # paid via DC doc #
            row(None, '2026-08-03 17:00.01', 'C400000001', 'U1111111AB',
                2500),
            row(None, '2026-08-03 17:00.01', 'C400000002', '22222222AI',
                200),                                   # credited at 150
            row(None, '2026-08-05 17:00.01', 'C400000003', '44444444AA',
                100),                                   # 61 days, unpaid
            row(None, '2026-08-06 17:00.01', 'C400000004', '55555555AA',
                125, req='2026-09-28'),                 # requested before
            row(None, '2026-08-07 17:00.01', 'C400000005', '66666666AA',
                90),                                    # 59 days
            row(None, '2026-10-01 17:00.01', 'C400000006', '77777777AA',
                40),                                    # recent
        ]
        self.master = os.path.join(self.tmp, 'Core_Returns_RECONCILED.xlsx')
        make_master(self.master, [dict(r) for r in self.rows])

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _rec(self, **opts):
        o = {'asof': self.asof, 'threshold': 60, 'dealer': '12345',
             'mark_requested': True, 'apply_dc': True,
             'include_prev': False}
        o.update(opts)
        rec = engine.Reconciler(o)
        rec.load_master(self.master, 'Core_Returns_RECONCILED.xlsx')
        rec.load_dc([], '')
        rec.pages = []
        return rec

    def test_credits(self):
        rec = self._rec()
        credits = parse_memo_text(MEMO_TEXT)['credits']
        for c in credits:
            rec._apply_credit(c, False)
        led = rec.ledger
        self.assertEqual(led.find('C400000001')[0]['memo'], 'CC00900001')
        # partial credit: master now shows the credited amount
        self.assertEqual(led.find('C400000002')[0]['amount'], 150)
        self.assertEqual(len(rec.r['partial']), 1)
        self.assertEqual(rec.r['partial'][0]['diff'], -50.0)
        # #000900001 on the master is the same memo as CC00900001
        self.assertEqual(rec.r['already_applied'], 1)
        self.assertEqual(len(led.find('C400000099')), 1)
        # applying the same memo again changes nothing
        before = len(led.rows)
        for c in credits:
            rec._apply_credit(c, False)
        self.assertEqual(len(led.rows), before)
        self.assertEqual(rec.r['already_applied'], 1 + 3)

    def test_unmatched_credit_is_added_as_paid(self):
        rec = self._rec()
        rec._apply_credit({'ticket': 'C499999999', 'memo': 'CC00900002',
                           'memo_date': self.asof, 'part': '99999999AA',
                           'claim': None, 'amount': 60.0}, False)
        added = rec.ledger.find('C499999999')[0]
        self.assertEqual((added['memo'], added['shipment'], added['amount']),
                         ('CC00900002', None, 60))
        self.assertEqual(len(rec.r['unmatched_credits']), 1)

    def test_credit_request_rules(self):
        rec = self._rec()
        rec.build_request()
        req = [x['ticket'] for x in rec.r['credit_request']]
        # 63 and 61 days old -> requested; 59 days -> not yet;
        # requested 9/28 -> listed separately, not repeated
        self.assertEqual(req, ['C400000001', 'C400000002', 'C400000003'])
        self.assertEqual([x['ticket'] for x in rec.r['prev_requested']],
                         ['C400000004'])
        self.assertEqual(rec.ledger.find('C400000003')[0]['requested'],
                         '2026-10-05')
        # oldest shipment first, no pages uploaded -> all need a re-scan
        self.assertEqual(rec.r['rescan'][0]['date'], '2026-08-03')

    def test_same_day_rerun_keeps_request(self):
        rec = self._rec()
        rec.build_request()
        rec2 = engine.Reconciler(rec.o)
        rec2.ledger = rec.ledger
        rec2.page_refs = {}
        rec2.build_request()
        self.assertEqual(len(rec2.r['credit_request']), 3)

    def test_include_previously_requested(self):
        rec = self._rec(include_prev=True)
        rec.build_request()
        self.assertIn('C400000004',
                      [x['ticket'] for x in rec.r['credit_request']])

    def test_dealerconnect(self):
        text = ('Tracking Number\tPart Number\tWarranty Claim\tAmount\t'
                'Received Date\tPayment Type\tDocument\n'
                'C400000003\t44444444AB\t\t100.00\t\t\t\n'      # part fix
                'C400000005\t66666666AA\t\t90.00\t9/30/2026\tCR\t#000900003\n'
                'C400000777\t88888888AA\t\t25.00\t\t\t\n')       # not on ours
        from recon.dealerconnect import merge
        rec = self._rec()
        rec.dc_rows = read_dc_text(text)
        rec.dc = merge(rec.dc_rows)
        rec.dc_has_paid = rec.dc_has_unpaid = True
        rec.dc_pull = datetime.date(2026, 10, 2)
        rec.apply_dc()
        status = {x['ticket']: x['status'] for x in rec.r['comparison']}
        self.assertEqual(status['C400000003'], 'FIXED FROM DEALERCONNECT')
        self.assertEqual(rec.ledger.find('C400000003')[0]['part'],
                         '44444444AB')
        self.assertEqual(status['C400000005'], 'PAID PER MOPAR')
        self.assertEqual(rec.ledger.find('C400000005')[0]['memo'],
                         '#000900003')
        self.assertEqual(status['C400000777'], 'ADD TO MASTER')
        self.assertEqual(status['C400000001'], 'MOPAR HAS NO RECORD')
        # shipped 10/01, pull dated 10/02 from the file name: 10/01 is
        # before the pull so it is checked; anything on/after is skipped
        self.assertIn('C400000006', status)


class TestScanDedupe(unittest.TestCase):
    def test_copies_vs_consecutive_tickets(self):
        def ln(t):
            return {'ticket': t, 'ticket_raw': t, 'ticket_fixed': False,
                    'claim': None, 'part': '11111111AA', 'part_ok': True,
                    'description': 'X', 'qty': 1, 'amount': 50.0,
                    'box': (10, 10, 20, 20), 'line_box': (0, 0, 1, 1),
                    'conf': 90}

        def page(f, n, sid, tickets):
            return {'file': f, 'path': f, 'page': n, 'width': 1700,
                    'height': 2200, 'dpi': 200, 'shipment_id': sid,
                    'dealer': None, 'lines': [ln(t) for t in tickets],
                    'checkoff_x': None, 'printed_total': None}
        rec = engine.Reconciler({'asof': datetime.date(2026, 10, 5),
                                 'threshold': 60})
        rec.load_master(None, None)
        rec.load_dc([], '')
        s1, s2 = '2026-10-01 17:00.01', '2026-10-03 17:00.01'
        rec.pages = [
            page('a.pdf', 1, s1, ['C412200212', 'C412200213', 'C412200300']),
            # second copy of the same page, 13 misread as 18
            page('b.pdf', 1, s1, ['C412200212', 'C412200218', 'C412200300']),
            # two-page shipper: consecutive tickets run onto page 2
            page('d.pdf', 1, s2, ['C413300001', 'C413300002']),
            page('d.pdf', 2, s2, ['C413300003', 'C413300004'])]
        rec.parts = engine.PartCorrector([])
        rec._resolve_lines()
        rec._apply_scans()
        added = [x['ticket'] for x in rec.r['new_shipments']]
        self.assertEqual(added, ['C412200212', 'C412200213', 'C412200300',
                                 'C413300001', 'C413300002', 'C413300003',
                                 'C413300004'])
        self.assertIn('C412200213', rec.r['ocr_review'][0]['issue'])


def scan_line(ticket, part='11111111AA', amount=50.0, y=0):
    return {'ticket': ticket, 'ticket_raw': ticket, 'ticket_fixed': False,
            'claim': None, 'part': part, 'part_ok': True,
            'description': 'X', 'qty': 1, 'amount': amount,
            'box': (10, y, 20, y + 10), 'line_box': (0, y, 1, y + 10),
            'conf': 90}


def scan_page(name, number, shipment, lines):
    for y, ln in enumerate(lines):
        ln['box'] = (10, 100 * y, 20, 100 * y + 10)
    return {'file': name, 'path': name, 'page': number, 'width': 1700,
            'height': 2200, 'dpi': 200, 'shipment_id': shipment,
            'dealer': None, 'lines': lines, 'checkoff_x': None,
            'printed_total': None}


class TestMisreadRules(unittest.TestCase):
    """GCRS control tickets run in sequence, so real tickets one digit
    apart sit next to each other on a page. A line is only ever treated
    as a misread of a neighbouring ticket when the paperwork does not show
    that neighbour under its own number and the part and amount agree."""

    SHP = '2026-07-01 17:00.01'

    def rec(self, master_rows=(), dc_text=''):
        from recon.dealerconnect import merge
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        rec = engine.Reconciler({'asof': datetime.date(2026, 10, 5),
                                 'threshold': 60})
        if master_rows:
            path = os.path.join(tmp, 'm.xlsx')
            make_master(path, [dict(r) for r in master_rows])
            rec.load_master(path, 'm.xlsx')
        else:
            rec.load_master(None, None)
        rec.load_dc([], dc_text)
        self.assertEqual(rec.dc, merge(rec.dc_rows))
        return rec

    def scan(self, rec, pages):
        rec.pages = pages
        rec.parts = engine.PartCorrector(
            [r.get('part') for r in rec.ledger.rows] +
            [d['part'] for d in rec.dc.values() if d.get('part')])
        rec._resolve_lines()
        rec._apply_scans()
        return [x['ticket'] for x in rec.r['new_shipments']]

    def test_a_neighbour_on_dealerconnect_never_replaces_a_ticket(self):
        # five tickets in a row on one page; Mopar's unpaid list has only
        # the second (the first is already paid)
        rec = self.rec(dc_text='C418800202\t22222222AB\t\t125.00\t\t\t')
        lines = [scan_line('C418800201', '11111111AA', 125.0),
                 scan_line('C418800202', '22222222AB', 125.0),
                 scan_line('C418800203', '33333333AC', 60.0),
                 scan_line('C418800204', '44444444AD', 75.0),
                 scan_line('C418800205', '55555555AE', 90.0)]
        added = self.scan(rec, [scan_page('a.pdf', 1, self.SHP, lines)])
        self.assertEqual(added, ['C418800201', 'C418800202', 'C418800203',
                                 'C418800204', 'C418800205'])
        self.assertEqual(rec.r['ocr_review'], [])
        by = {x['ticket']: x for x in rec.r['new_shipments']}
        self.assertEqual(by['C418800201']['source'], 'scan')
        self.assertEqual(by['C418800202']['source'], 'DealerCONNECT')
        # each ticket points at its own line (where "UnPaid" is stamped)
        for i, t in enumerate(added):
            page, ln = rec.page_refs[t][0]
            self.assertEqual(ln['box'][1], 100 * i, t)
        # ... also when the DealerCONNECT ticket is on the next page
        rec = self.rec(dc_text='C418800202\t22222222AB\t\t125.00\t\t\t')
        added = self.scan(rec, [
            scan_page('a.pdf', 1, self.SHP,
                      [scan_line('C418800150', '99999999AA', 10.0),
                       scan_line('C418800201', '11111111AA', 125.0)]),
            scan_page('a.pdf', 2, self.SHP,
                      [scan_line('C418800202', '22222222AB', 125.0),
                       scan_line('C418800203', '33333333AC', 60.0)])])
        self.assertEqual(added, ['C418800150', 'C418800201', 'C418800202',
                                 'C418800203'])

    def test_a_real_misread_is_still_fixed_from_dealerconnect(self):
        # ...207 scanned as ...287; DealerCONNECT has ...207 with the same
        # part and amount, and it is on no page under its own number
        rec = self.rec(dc_text='C418800207\t22222222AB\t\t125.00\t\t\t')
        added = self.scan(rec, [scan_page('a.pdf', 1, self.SHP, [
            scan_line('C418800201', '11111111AA', 50.0),
            scan_line('C418800287', '22222222AB', 125.0)])])
        self.assertEqual(added, ['C418800201', 'C418800207'])
        # (...201 is one digit from ...207 too, but its part and amount
        # are different, so it is left alone without comment)
        self.assertEqual(len(rec.r['ocr_review']), 1)
        self.assertIn('C418800207', rec.r['ocr_review'][0]['issue'])
        self.assertIn('used that', rec.r['ocr_review'][0]['issue'])

    def test_a_neighbour_with_another_part_or_amount_is_another_ticket(self):
        # DealerCONNECT's ...202 is on no uploaded page, but the line that
        # reads ...201 shows a different part and amount: it is ...201
        rec = self.rec(dc_text='C418800202\t22222222AB\t\t125.00\t\t\t')
        added = self.scan(rec, [scan_page('a.pdf', 1, self.SHP, [
            scan_line('C418800201', '11111111AA', 50.0)])])
        self.assertEqual(added, ['C418800201'])
        self.assertEqual(rec.ledger.find('C418800202'), [])
        self.assertEqual(rec.r['ocr_review'], [])
        # the same part with a different amount is not enough either
        rec = self.rec(dc_text='C418800202\t11111111AA\t\t125.00\t\t\t')
        added = self.scan(rec, [scan_page('a.pdf', 1, self.SHP, [
            scan_line('C418800201', '11111111AA', 50.0)])])
        self.assertEqual(added, ['C418800201'])
        self.assertEqual(rec.r['ocr_review'], [])
        # nothing readable on the line to compare: kept as read, and said
        rec = self.rec(dc_text='C418800202\t22222222AB\t\t125.00\t\t\t')
        added = self.scan(rec, [scan_page('a.pdf', 1, self.SHP, [
            scan_line('C418800201', '', None)])])
        self.assertEqual(added, ['C418800201'])
        note = rec.r['ocr_review'][0]
        self.assertEqual(note['ticket'], 'C418800201')
        self.assertIn('Kept as C418800201', note['issue'])

    def test_master_neighbour_on_another_page_is_not_a_misread(self):
        master = [row(None, self.SHP, 'C418800202', '22222222AB', 125)]
        rec = self.rec(master)
        added = self.scan(rec, [
            scan_page('a.pdf', 1, self.SHP,
                      [scan_line('C418800150', '99999999AA', 10.0),
                       scan_line('C418800201', '22222222AB', 125.0)]),
            scan_page('a.pdf', 2, self.SHP,
                      [scan_line('C418800202', '22222222AB', 125.0),
                       scan_line('C418800203', '33333333AC', 60.0)])])
        # ...201 is its own ticket even with the same part and amount,
        # because ...202 is right there on page 2
        self.assertEqual(added, ['C418800150', 'C418800201', 'C418800203'])
        self.assertEqual(rec.r['ocr_review'], [])
        self.assertEqual(len(rec.page_refs['C418800202']), 1)

    def test_master_neighbour_is_matched_only_when_the_line_agrees(self):
        master = [row(None, self.SHP, 'C418800202', '22222222AB', 125),
                  row(None, self.SHP, 'C418800150', '99999999AA', 10)]
        # agrees (same part and amount), and ...202 is on no page: misread
        rec = self.rec(master)
        added = self.scan(rec, [scan_page('a.pdf', 1, self.SHP, [
            scan_line('C418800150', '99999999AA', 10.0),
            scan_line('C418800282', '22222222AB', 125.0)])])
        self.assertEqual(added, [])
        self.assertIn('matched to C418800202',
                      rec.r['ocr_review'][0]['issue'])
        self.assertEqual(len(rec.page_refs['C418800202']), 1)
        # a different amount: kept as its own ticket, and flagged
        rec = self.rec(master)
        added = self.scan(rec, [scan_page('a.pdf', 1, self.SHP, [
            scan_line('C418800150', '99999999AA', 10.0),
            scan_line('C418800201', '22222222AB', 60.0)])])
        self.assertEqual(added, ['C418800201'])
        self.assertIn('kept as its own ticket',
                      rec.r['ocr_review'][0]['issue'])
        self.assertEqual(rec.page_refs.get('C418800202', []), [])

    def test_second_copy_of_a_page_with_a_misread_adds_nothing(self):
        s1 = '2026-10-01 17:00.01'
        master = [row(None, s1, t, '11111111AA', 50) for t in
                  ('C412200212', 'C412200213', 'C412200300')]
        rec = self.rec(master)
        added = self.scan(rec, [
            scan_page('a.pdf', 1, s1, [scan_line(t) for t in (
                'C412200212', 'C412200213', 'C412200300')]),
            scan_page('b.pdf', 1, s1, [scan_line(t) for t in (
                'C412200212', 'C412200218', 'C412200300')])])
        self.assertEqual(added, [])                 # ...218 is ...213
        self.assertEqual(len(rec.page_refs['C412200213']), 2)
        # the same with new tickets that DealerCONNECT knows, misread
        # copy first
        dc = '\n'.join(f'{t}\t11111111AA\t\t50.00\t\t\t' for t in (
            'C412200212', 'C412200213', 'C412200300'))
        rec = self.rec(dc_text=dc)
        added = self.scan(rec, [
            scan_page('b.pdf', 1, s1, [scan_line(t) for t in (
                'C412200212', 'C412200218', 'C412200300')]),
            scan_page('a.pdf', 1, s1, [scan_line(t) for t in (
                'C412200212', 'C412200213', 'C412200300')])])
        self.assertEqual(added, ['C412200212', 'C412200213', 'C412200300'])

    def test_scanned_memo_credit_goes_to_a_neighbour_only_when_sure(self):
        def credit(ticket, amount, memo='CC00900009'):
            return {'ticket': ticket, 'memo': memo,
                    'memo_date': datetime.date(2026, 10, 3),
                    'part': None, 'claim': None, 'amount': amount}
        master = [row(None, self.SHP, 'C418800202', '22222222AB', 125)]
        # same amount, and ...202 has no credit of its own: applied, said
        rec = self.rec(master)
        rec._apply_credit(credit('C418800282', 125.0), True)
        self.assertEqual(rec.ledger.find('C418800202')[0]['memo'],
                         'CC00900009')
        self.assertIn('applied to C418800202', rec.r['warnings'][0])
        # a different amount: its own (unmatched) credit, ...202 stays open
        rec = self.rec(master)
        rec._apply_credit(credit('C418800201', 60.0), True)
        self.assertIsNone(rec.ledger.find('C418800202')[0]['memo'])
        self.assertEqual(len(rec.r['unmatched_credits']), 1)
        # ...202 is credited in the same batch: ...201 is not a misread
        rec = self.rec(master)
        rec.credit_tickets = {'C418800201', 'C418800202'}
        rec._apply_credit(credit('C418800201', 125.0), True)
        rec._apply_credit(credit('C418800202', 125.0, 'CC00900010'), True)
        self.assertEqual(rec.ledger.find('C418800202')[0]['memo'],
                         'CC00900010')
        self.assertEqual(rec.ledger.find('C418800201')[0]['memo'],
                         'CC00900009')

    def test_a_request_made_in_this_run_shows_on_the_unpaid_list(self):
        rec = self.rec([row(None, self.SHP, 'C418800202', '22222222AB',
                            125)])
        rec.pages = []
        rec.build_request()
        self.assertEqual(rec.r['unpaid'][0]['requested'], '2026-10-05')
        rec = self.rec([row(None, self.SHP, 'C418800202', '22222222AB',
                            125)])
        rec.o['mark_requested'] = False
        rec.pages = []
        rec.build_request()
        self.assertIsNone(rec.r['unpaid'][0]['requested'])


class TestFilesThatGiveNothing(unittest.TestCase):
    """A file in the wrong box, or one that cannot be read, is called out
    instead of quietly producing zeros."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def run_engine(self, **inputs):
        base = {'master': None, 'memos': [], 'shippers': [], 'dc': [],
                'dc_text': ''}
        base.update(inputs)
        out = os.path.join(self.tmp, 'out')
        os.makedirs(out, exist_ok=True)
        return engine.run(base, {'asof': datetime.date(2026, 10, 5),
                                 'threshold': 60}, out)

    def write(self, name, data):
        path = os.path.join(self.tmp, name)
        with open(path, 'wb') as fh:
            fh.write(data)
        return path, name

    def test_shipper_file_with_no_tickets_is_called_out(self):
        for name, data in (('empty.pdf', b''),
                           ('broken.pdf', b'%PDF-1.4 this is not a pdf'),
                           ('letter.pdf', b'PK\x03\x04 a word file')):
            res = self.run_engine(shippers=[self.write(name, data)])
            warn = [w for w in res['warnings'] if w.startswith(name)]
            self.assertEqual(len(warn), 1, res['warnings'])
            self.assertIn('no control tickets could be read', warn[0])
            self.assertNotIn('Traceback', warn[0])

    def test_dealerconnect_file_that_is_not_excel_is_left_out(self):
        res = self.run_engine(
            dc=[self.write('unpaid.xlsx', b'')],
            dc_text='C400000001\t11111111AA\t\t50.00\t\t\t')
        warn = [w for w in res['warnings'] if w.startswith('unpaid.xlsx')]
        self.assertEqual(len(warn), 1, res['warnings'])
        self.assertIn('could not be opened', warn[0])
        # the pasted rows were still used
        self.assertEqual(res['inputs']['dealerconnect'],
                         ['unpaid.xlsx', 'pasted rows'])

    def test_memo_that_cannot_be_read_is_called_out(self):
        res = self.run_engine(memos=[self.write('memo.pdf', b'')])
        self.assertTrue(any(w.startswith('memo.pdf') for w in
                            res['warnings']), res['warnings'])


class TestMasterWorkbook(unittest.TestCase):
    def test_round_trip_and_layout(self):
        tmp = tempfile.mkdtemp()
        try:
            p = os.path.join(tmp, 'm.xlsx')
            rows = [row(None, '2026-08-03 17:00.01', 'C400000001', 'A1', 5),
                    row('CC1', None, 'C400000002', 'B1', 7),
                    row(None, '2026-07-01 17:00.01', 'C400000003', 'C1', 9)]
            M.write_master(rows, p)
            back, info = M.read_master(p)
            # blank Shipment ID first, then oldest shipment first
            self.assertEqual([r['ticket'] for r in back],
                             ['C400000002', 'C400000003', 'C400000001'])
            wb = load_workbook(p)
            self.assertEqual(wb.sheetnames, ['Core Returns', 'Unpaid Cores'])
            ws = wb['Core Returns']
            self.assertEqual([c.value for c in ws[1]][:8], M.HEADERS)
            self.assertEqual(ws['J1'].value, '=TODAY()')
            self.assertTrue(ws['G2'].value.startswith('=IF(B2=""'))
            self.assertEqual(ws.freeze_panes, 'A2')
            self.assertEqual(
                [c.value for c in wb['Unpaid Cores']['C']][1:],
                ['C400000003', 'C400000001'])
            cf = list(ws.conditional_formatting)
            self.assertEqual(cf[0].rules[0].formula,
                             ['AND($A2="",ISNUMBER($G2),$G2>60)'])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestWebApp(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        os.environ['RECON_JOB_DIR'] = self.tmp
        from recon import jobs
        jobs.JOB_ROOT = self.tmp
        import app as webapp
        self.client = webapp.app.test_client()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_pages(self):
        self.assertEqual(self.client.get('/').status_code, 200)
        self.assertEqual(self.client.get('/reconcile').status_code, 200)

    def test_needs_input(self):
        r = self.client.post('/api/reconcile', data={})
        self.assertEqual(r.status_code, 400)

    def test_wrong_file_type(self):
        import io
        r = self.client.post('/api/reconcile', data={
            'master': (io.BytesIO(b'x'), 'master.pdf')},
            content_type='multipart/form-data')
        self.assertEqual(r.status_code, 400)

    def test_full_run_with_pasted_dealerconnect(self):
        import io
        p = os.path.join(self.tmp, 'Core_Returns_RECONCILED.xlsx')
        M.write_master([row(None, '2026-07-01 17:00.01', 'C400000001',
                            '11111111AA', 50)], p)
        with open(p, 'rb') as fh:
            data = {'master': (io.BytesIO(fh.read()),
                               'Core_Returns_RECONCILED.xlsx'),
                    'dc_text': 'C400000001\t11111111AA\t\t50.00\t\t\t',
                    'asof': '2026-10-05', 'threshold': '60'}
        r = self.client.post('/api/reconcile', data=data,
                             content_type='multipart/form-data')
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        job_id = r.get_json()['job_id']
        for _ in range(100):
            job = self.client.get(f'/api/jobs/{job_id}').get_json()
            if job['status'] in ('done', 'error'):
                break
            time.sleep(0.1)
        self.assertEqual(job['status'], 'done', job.get('error'))
        res = job['result']
        self.assertEqual([x['ticket'] for x in res['credit_request']],
                         ['C400000001'])
        names = [f['name'] for f in res['files']]
        self.assertIn('Credit_Request.xlsx', names)
        dl = self.client.get(f'/api/jobs/{job_id}/files/Credit_Request.xlsx')
        self.assertEqual(dl.status_code, 200)
        wb = load_workbook(io.BytesIO(dl.data))
        ws = wb['Credit Request']
        self.assertEqual([c.value for c in ws[2]],
                         [None, '2026-07-01', 'C400000001'])
        self.assertEqual(
            self.client.get(f'/api/jobs/{job_id}/files/../app.py')
            .status_code, 404)


if __name__ == '__main__':
    unittest.main()
