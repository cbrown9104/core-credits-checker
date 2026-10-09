"""Tests for store accounts: sign-in links, one store never seeing another's
data, the saved master and the run history.

These need a Postgres database to run against. Point TEST_DATABASE_URL at
an EMPTY scratch database (everything in it is dropped):

    TEST_DATABASE_URL=postgresql://user@localhost/scratch \\
        python -m unittest discover -s tests -v

Without TEST_DATABASE_URL these tests are skipped and the rest still run.
All data is made up.
"""
import io
import os
import re
import shutil
import sys
import tempfile
import time
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(
    __file__))))

TEST_DB = os.environ.get('TEST_DATABASE_URL', '').strip()
OWNER = 'owner@example.com'


def row(memo, shp, ticket, part, amount, claim=None, req=None):
    return {'memo': memo, 'shipment': shp, 'ticket': ticket, 'part': part,
            'claim': claim, 'amount': amount, 'requested': req}


def master_bytes(rows):
    from recon import master as M
    d = tempfile.mkdtemp()
    try:
        p = os.path.join(d, 'm.xlsx')
        M.write_master(rows, p)
        with open(p, 'rb') as fh:
            return fh.read()
    finally:
        shutil.rmtree(d, ignore_errors=True)


def many_rows(n, start=0):
    return [row(None, '2026-09-01 17:00.01', f'C4100{start + i:05d}',
                '11111111AA', 50) for i in range(n)]


class TestOpenMode(unittest.TestCase):
    """With no database configured nothing about the app changes."""

    def setUp(self):
        self.saved = os.environ.pop('DATABASE_URL', None)
        import app as webapp
        self.client = webapp.app.test_client()

    def tearDown(self):
        if self.saved is not None:
            os.environ['DATABASE_URL'] = self.saved

    def test_pages_are_open_and_account_pages_do_not_exist(self):
        home = self.client.get('/')
        self.assertEqual(home.status_code, 200)
        self.assertIn(b'Quick Check', home.data)
        self.assertNotIn(b'Sign out', home.data)
        self.assertEqual(self.client.get('/reconcile').status_code, 200)
        for path in ('/login', '/history', '/users', '/owner', '/stores',
                     '/master/download'):
            self.assertEqual(self.client.get(path).status_code, 404, path)
        # posting needs no token when accounts are off
        self.assertEqual(self.client.post('/api/reconcile',
                                          data={}).status_code, 400)
        self.assertNotIn('Set-Cookie', home.headers)

    def test_memo_number_prefix_is_not_fixed(self):
        from app import parse_credits
        for prefix in ('03181000', '04220000', ''):
            text = (f'CREDIT MEMO NUMBER: {prefix}CC00900001\n'
                    f'REFERENCE/CONTROL NUMBER C400000001\n')
            self.assertEqual(parse_credits(text),
                             {'C400000001': 'CC00900001'}, prefix)


@unittest.skipUnless(TEST_DB, 'set TEST_DATABASE_URL to run account tests')
class AccountsCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.env = {k: os.environ.get(k) for k in (
            'DATABASE_URL', 'MAIL_BACKEND', 'OWNER_EMAIL', 'APP_BASE_URL',
            'RENDER_EXTERNAL_URL', 'RECON_JOB_DIR')}
        cls.tmp = tempfile.mkdtemp()
        os.environ.pop('RENDER_EXTERNAL_URL', None)
        os.environ.update(DATABASE_URL=TEST_DB, MAIL_BACKEND='memory',
                          OWNER_EMAIL=OWNER, RECON_JOB_DIR=cls.tmp,
                          APP_BASE_URL='http://localhost')
        from accounts import db
        db.close()
        import psycopg
        with psycopg.connect(TEST_DB, autocommit=True) as conn:
            conn.execute('DROP SCHEMA public CASCADE')
            conn.execute('CREATE SCHEMA public')
        from recon import jobs
        jobs.JOB_ROOT = cls.tmp
        import app as webapp
        cls.app = webapp.app

    @classmethod
    def tearDownClass(cls):
        from accounts import db
        db.close()
        for k, v in cls.env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def setUp(self):
        from accounts import db, mailer
        mailer.outbox.clear()
        # each test starts with no sign-in requests counted against it
        db.execute('DELETE FROM login_attempts')

    # ── helpers ─────────────────────────────────────────────────────────
    def client(self):
        return self.app.test_client()

    def guard(self, c):
        """What a freshly loaded page hands the browser: (page token,
        id of the store the page is showing)."""
        c.get('/login', follow_redirects=True)
        with c.session_transaction() as s:
            return s['csrf'], str(s.get('sid', ''))

    def post(self, c, path, data=None, **kw):
        data = dict(data or {})
        token, store = self.guard(c)
        data.setdefault('csrf_token', token)
        data.setdefault('store_check', store)
        return c.post(path, data=data, **kw)

    def last_link(self, to):
        from accounts import mailer
        mails = [m for m in mailer.outbox if m['to'] == to]
        self.assertTrue(mails, f'no mail to {to}')
        m = re.search(r'http://localhost(/login/link/\S+)',
                      mails[-1]['text'])
        self.assertTrue(m, mails[-1]['text'])
        return m.group(1)

    def use_link(self, c, link):
        self.assertEqual(c.get(link).status_code, 200)
        r = self.post(c, link)
        self.assertEqual(r.status_code, 302, r.get_data(as_text=True))
        return r

    def sign_in(self, email):
        c = self.client()
        r = self.post(c, '/login', {'email': email})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        self.use_link(c, self.last_link(email))
        return c

    def make_store(self, name, admin, dealer='12345'):
        owner = self.sign_in(OWNER)
        r = self.post(owner, '/owner/stores', {
            'name': name, 'dealer_code': dealer, 'admin_email': admin,
            'send_invite': '1'})
        self.assertEqual(r.status_code, 302)
        from accounts import db
        return db.one('SELECT * FROM stores WHERE name = %s ORDER BY id '
                      'DESC LIMIT 1', (name,))

    def start_recon(self, c, data):
        token, store = self.guard(c)
        return c.post('/api/reconcile', data=data,
                      content_type='multipart/form-data',
                      headers={'X-CSRF-Token': token, 'X-Store-Id': store})

    def run_recon(self, c, data):
        r = self.start_recon(c, data)
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        job_id = r.get_json()['job_id']
        for _ in range(400):
            job = c.get(f'/api/jobs/{job_id}').get_json()
            if job['status'] in ('done', 'error'):
                break
            time.sleep(0.05)
        return job_id, job

    def first_run(self, c, ticket='C400000001'):
        content = master_bytes([row(None, '2026-07-01 17:00.01', ticket,
                                    '11111111AA', 50)])
        return self.run_recon(c, {
            'master': (io.BytesIO(content), 'Core_Returns_RECONCILED.xlsx'),
            'dc_text': f'{ticket}\t11111111AA\t\t50.00\t\t\t',
            'asof': '2026-10-05', 'threshold': '60'})


class TestSignIn(AccountsCase):
    def test_signed_out_visitors_see_the_front_page_only(self):
        c = self.client()
        home = c.get('/')
        self.assertEqual(home.status_code, 200)
        self.assertIn(b'Sign in', home.data)
        self.assertNotIn(b'Drop shipper PDF', home.data)
        r = c.get('/reconcile')
        self.assertEqual(r.status_code, 302)
        self.assertIn('/login', r.headers['Location'])
        for path in ('/history', '/users', '/owner', '/stores',
                     '/master/download'):
            self.assertEqual(c.get(path).status_code, 302, path)
        self.assertEqual(c.get('/api/jobs/' + '0' * 32).status_code, 401)
        # a signed-out upload is turned away without a page token check
        # (so before its body is read)
        r = c.post('/api/reconcile', data={
            'memos': (io.BytesIO(b'x' * 200000), 'big.pdf')},
            content_type='multipart/form-data')
        self.assertEqual(r.status_code, 401)
        self.assertEqual(c.post('/api/check-cores').status_code, 401)
        # a sign-in post must say how big it is, and be small
        token, _ = self.guard(c)
        form = {'email': 'x@nowhere.test', 'csrf_token': token}
        r = c.post('/login', data=form, environ_overrides={
            'CONTENT_LENGTH': '', 'wsgi.input_terminated': True})
        self.assertEqual(r.status_code, 411)        # a streamed upload
        r = c.post('/login', data=form, environ_overrides={
            'CONTENT_LENGTH': ''}, headers={'Transfer-Encoding': 'chunked'})
        self.assertEqual(r.status_code, 411)
        r = c.post('/login', data=dict(form, pad='x' * 70000))
        self.assertEqual(r.status_code, 413)

    def test_link_signs_in_once(self):
        self.make_store('Link Store', 'pat@linkstore.test')
        c = self.client()
        self.post(c, '/login', {'email': 'Pat@LinkStore.test '})
        link = self.last_link('pat@linkstore.test')
        # opening the link (as a mail scanner would) does not use it up
        self.assertEqual(self.client().get(link).status_code, 200)
        self.use_link(c, link)
        page = c.get('/reconcile')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'Link Store', page.data)
        self.assertIn(b'pat@linkstore.test', page.data)
        # the same link is dead now, for everyone
        other = self.client()
        self.assertEqual(other.get(link).status_code, 410)
        self.assertEqual(self.post(other, link).status_code, 410)

    def test_a_scanner_probing_the_link_neither_uses_it_nor_signs_in(self):
        self.make_store('Probe Store', 'pru@probe.test')
        c = self.client()
        self.post(c, '/login', {'email': 'pru@probe.test'})
        link = self.last_link('pru@probe.test')
        probe = self.client()
        for _ in range(3):
            self.assertEqual(probe.head(link).status_code, 200)
            self.assertEqual(probe.open(link, method='OPTIONS').status_code,
                             200)
            with probe.session_transaction() as s:
                self.assertNotIn('sx', s)           # not signed in
        self.assertEqual(probe.get('/reconcile').status_code, 302)
        # a post without the page's token does nothing either
        self.assertEqual(probe.post(link).status_code, 400)
        # the real person can still use it
        self.use_link(c, link)
        self.assertEqual(c.get('/reconcile').status_code, 200)
        # HEAD on the sign-in form never sends mail
        from accounts import mailer
        mailer.outbox.clear()
        self.client().head('/login', data={'email': 'pru@probe.test'})
        self.assertEqual(mailer.outbox, [])

    def test_using_one_link_retires_the_others(self):
        self.make_store('Retire Store', 'rex@retire.test')
        invite = self.last_link('rex@retire.test')
        c = self.client()
        self.post(c, '/login', {'email': 'rex@retire.test'})
        older = self.last_link('rex@retire.test')
        self.post(c, '/login', {'email': 'rex@retire.test'})
        newest = self.last_link('rex@retire.test')
        self.assertNotEqual(older, newest)
        self.use_link(c, newest)
        for dead in (older, invite):
            self.assertEqual(self.client().get(dead).status_code, 410)

    def test_unknown_email_gets_the_same_answer_and_no_mail(self):
        from accounts import mailer
        self.make_store('Known Store', 'kim@known.test')
        mailer.outbox.clear()
        known = self.post(self.client(), '/login',
                          {'email': 'kim@known.test'})
        unknown = self.post(self.client(), '/login',
                            {'email': 'nobody@nowhere.test'})
        self.assertEqual((known.status_code, unknown.status_code),
                         (200, 200))
        self.assertIn(b'Check your email', unknown.data)
        self.assertEqual(
            known.get_data(as_text=True).replace('kim@known.test', 'E'),
            unknown.get_data(as_text=True).replace('nobody@nowhere.test',
                                                    'E'))
        self.assertEqual([m['to'] for m in mailer.outbox],
                         ['kim@known.test'])
        for bad in ('', 'not-an-email', 'a@b', 'x\x00y@nowhere.test',
                    'a b@nowhere.test'):
            self.assertEqual(self.post(self.client(), '/login',
                                       {'email': bad}).status_code, 400, bad)

    def test_next_only_ever_leads_to_the_apps_own_pages(self):
        from accounts import web
        for good in ('/', '/reconcile', '/history', '/users', '/owner',
                     '/stores', '/history?x=1'):
            self.assertEqual(web._safe_next(good), good)
        for bad in ('/login/link/abc', '/login', '/logout', '/api/jobs/x',
                    '/switch-store', '/owner/open/3', '//evil.example',
                    'https://evil.example/', '/\\evil.example',
                    '/history\r\nX: y', 'reconcile', '', None,
                    '/master/1/restore'):
            self.assertEqual(web._safe_next(bad), '', bad)
        # end to end: a crafted sign-in address cannot chain the person
        # into someone else's sign-in link
        self.make_store('Next Store', 'nan@next.test')
        self.make_store('Trap Store', 'mal@trap.test')
        trap = self.last_link('mal@trap.test')
        c = self.client()
        page = c.get('/login?next=' + trap)
        self.assertNotIn(trap.encode(), page.data)
        self.post(c, '/login', {'email': 'nan@next.test', 'next': trap})
        r = self.use_link(c, self.last_link('nan@next.test'))
        self.assertTrue(r.headers['Location'].endswith('/reconcile'))
        # a real next is honoured once, and not reused by a later sign-in
        c2 = self.client()
        self.post(c2, '/login', {'email': 'nan@next.test',
                                 'next': '/history'})
        r = self.use_link(c2, self.last_link('nan@next.test'))
        self.assertTrue(r.headers['Location'].endswith('/history'))
        self.post(c2, '/logout')
        self.post(c2, '/login', {'email': 'nan@next.test'})
        r = self.use_link(c2, self.last_link('nan@next.test'))
        self.assertTrue(r.headers['Location'].endswith('/reconcile'))
        # opening another person's link while signed in says so plainly
        warn = c.get(trap).get_data(as_text=True)
        self.assertIn('signed in as', warn)
        self.assertIn('nan@next.test', warn)
        self.assertIn('mal@trap.test', warn)

    def test_expired_link_is_dead(self):
        from accounts import core, db
        self.make_store('Expiry Store', 'ed@expiry.test')
        user = core.user_by_email('ed@expiry.test')
        token = core.new_login_token(user['id'])
        db.execute("UPDATE login_tokens SET expires_at = now() - "
                   "interval '1 minute' WHERE user_id = %s", (user['id'],))
        c = self.client()
        self.assertEqual(c.get(f'/login/link/{token}').status_code, 410)
        self.assertEqual(self.post(c, f'/login/link/{token}').status_code,
                         410)

    def test_too_many_requests_are_refused(self):
        from accounts import core
        c = self.client()
        codes = [self.post(c, '/login',
                           {'email': 'flood@nowhere.test'}).status_code
                 for _ in range(core.MAX_PER_EMAIL_AND_IP + 1)]
        self.assertEqual(codes[:-1], [200] * core.MAX_PER_EMAIL_AND_IP)
        self.assertEqual(codes[-1], 429)
        # a refused flood does not run up the email's own allowance:
        # the real person, asking from their own address, still gets in
        for _ in range(40):
            self.post(c, '/login', {'email': 'flood@nowhere.test'})
        token, _ = self.guard(c)
        r = c.post('/login', data={'email': 'flood@nowhere.test',
                                   'csrf_token': token},
                   headers={'CF-Connecting-IP': '203.0.113.9'})
        self.assertEqual(r.status_code, 200)

    def test_posts_need_the_page_token(self):
        c = self.sign_in(OWNER)
        self.assertEqual(c.post('/logout').status_code, 400)
        self.assertEqual(c.post('/api/reconcile', data={}).status_code, 400)
        self.assertEqual(c.post('/owner/stores', data={
            'name': 'x', 'admin_email': 'x@x.test',
            'csrf_token': 'wrong'}).status_code, 400)
        # still signed in: a refused post does not sign anyone out
        self.assertEqual(c.get('/owner').status_code, 200)

    def test_sign_out_ends_the_session_on_the_server(self):
        self.make_store('Out Store', 'oz@out.test')
        c = self.sign_in('oz@out.test')
        self.assertEqual(c.get('/reconcile').status_code, 200)
        stolen = c.get_cookie('pms_session').value
        self.assertEqual(self.post(c, '/logout').status_code, 302)
        self.assertEqual(c.get('/reconcile').status_code, 302)
        # a copy of the old cookie is worthless now
        thief = self.client()
        thief.set_cookie('pms_session', stolen)
        self.assertEqual(thief.get('/reconcile').status_code, 302)
        self.assertEqual(thief.get('/api/jobs/' + '0' * 32).status_code,
                         401)

    def test_the_cookie_is_not_resent_on_every_answer(self):
        self.make_store('Quiet Store', 'quy@quiet.test')
        c = self.sign_in('quy@quiet.test')
        c.get('/reconcile')
        r = c.get('/api/jobs/' + '0' * 32)        # like the status poll
        self.assertEqual(r.status_code, 404)
        self.assertNotIn('Set-Cookie', r.headers)

    def test_emailed_links_ignore_a_forged_host_header(self):
        from accounts import mailer
        self.make_store('Host Store', 'hal@host.test')
        mailer.outbox.clear()
        c = self.client()
        evil = 'http://evil.example'     # the whole exchange is forged
        c.get('/login', base_url=evil)
        with c.session_transaction(base_url=evil) as s:
            token = s['csrf']
        r = c.post('/login', data={'email': 'hal@host.test',
                                   'csrf_token': token}, base_url=evil)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(mailer.outbox), 1)
        text = mailer.outbox[0]['text']
        self.assertIn('http://localhost/login/link/', text)
        self.assertNotIn('evil.example', text)

    def test_no_configured_address_means_no_links_in_production(self):
        from accounts import mailer
        self.make_store('Base Store', 'bo@base.test')
        saved = {k: os.environ.pop(k, None)
                 for k in ('APP_BASE_URL', 'MAIL_BACKEND')}
        os.environ['RESEND_API_KEY'] = 'test-key-never-used'
        try:
            mailer.outbox.clear()
            r = self.post(self.client(), '/login', {'email': 'bo@base.test'})
            self.assertEqual(r.status_code, 503)
        finally:
            os.environ.pop('RESEND_API_KEY', None)
            for k, v in saved.items():
                if v is not None:
                    os.environ[k] = v


class TestStores(AccountsCase):
    def test_owner_pages_are_for_the_owner_only(self):
        self.make_store('Owner Test Store', 'amy@ots.test')
        amy = self.sign_in('amy@ots.test')
        self.assertEqual(amy.get('/owner').status_code, 404)
        self.assertEqual(self.post(amy, '/owner/stores', {
            'name': 'Sneaky', 'admin_email': 'amy@ots.test'}).status_code,
            404)
        owner = self.sign_in(OWNER)
        page = owner.get('/owner')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'Owner Test Store', page.data)
        # odd characters in a name are cleaned, never a server error
        r = self.post(owner, '/owner/stores', {
            'name': 'Odd\x00Name\tStore', 'admin_email': 'odd@odd.test'})
        self.assertEqual(r.status_code, 302)
        self.assertIn(b'Odd Name Store', owner.get('/owner').data)

    def test_users_page_rules(self):
        from accounts import core, mailer
        store = self.make_store('Users Store', 'ann@users.test')
        ann = self.sign_in('ann@users.test')
        mailer.outbox.clear()
        r = self.post(ann, '/users/add', {'email': 'bob@users.test',
                                          'role': 'user'})
        self.assertEqual(r.status_code, 302)
        self.assertEqual([m['to'] for m in mailer.outbox],
                         ['bob@users.test'])
        self.assertIn('Users Store', mailer.outbox[0]['subject'])
        bob_user = core.user_by_email('bob@users.test')
        # the invitation link signs Bob in, straight into his one store
        bob = self.client()
        self.use_link(bob, self.last_link('bob@users.test'))
        self.assertEqual(bob.get('/reconcile').status_code, 200)
        # a plain user cannot manage people
        self.assertEqual(bob.get('/users').status_code, 403)
        self.assertEqual(self.post(bob, '/users/add', {
            'email': 'eve@users.test', 'role': 'admin'}).status_code, 403)
        ann_user = core.user_by_email('ann@users.test')
        self.assertEqual(self.post(
            bob, f'/users/{ann_user["id"]}/remove').status_code, 403)
        self.assertEqual(self.post(bob, '/store/settings', {
            'name': 'Taken Over', 'threshold_days': '1'}).status_code, 403)
        # the only admin cannot be removed or demoted
        self.post(ann, f'/users/{ann_user["id"]}/remove')
        self.post(ann, '/users/add', {'email': 'ann@users.test',
                                      'role': 'user'})
        self.assertEqual(core.role_in(ann_user['id'], store['id']), 'admin')
        # an admin sets the threshold but not the store's name or code
        self.post(ann, '/store/settings', {
            'name': 'Renamed', 'dealer_code': '99999',
            'threshold_days': '45'})
        now = core.store_by_id(store['id'])
        self.assertEqual((now['name'], now['dealer_code'],
                          now['threshold_days']), ('Users Store', '12345',
                                                   45))
        # removing Bob signs him out on his very next click
        stolen = bob.get_cookie('pms_session').value
        self.assertEqual(self.post(
            ann, f'/users/{bob_user["id"]}/remove').status_code, 302)
        self.assertIsNone(core.role_in(bob_user['id'], store['id']))
        self.assertEqual(bob.get('/reconcile').status_code, 302)
        self.assertEqual(bob.get('/api/jobs/' + '0' * 32).status_code, 401)
        # and he cannot get a new link
        mailer.outbox.clear()
        self.post(self.client(), '/login', {'email': 'bob@users.test'})
        self.assertEqual(mailer.outbox, [])
        # added back later: his old browser session stays dead
        self.post(ann, '/users/add', {'email': 'bob@users.test',
                                      'role': 'user'})
        old = self.client()
        old.set_cookie('pms_session', stolen)
        self.assertEqual(old.get('/reconcile').status_code, 302)

    def test_link_sent_before_removal_stops_working(self):
        from accounts import core
        self.make_store('Removal Store', 'rae@removal.test')
        rae = self.sign_in('rae@removal.test')
        self.post(rae, '/users/add', {'email': 'sam@removal.test',
                                      'role': 'user'})
        link = self.last_link('sam@removal.test')
        sam = core.user_by_email('sam@removal.test')
        self.post(rae, f'/users/{sam["id"]}/remove')
        c = self.client()
        self.assertEqual(c.get(link).status_code, 410)
        self.assertEqual(self.post(c, link).status_code, 410)

    def test_one_person_two_stores_always_picks_for_themselves(self):
        a = self.make_store('Group Store A', 'gm@group.test', '11111')
        b = self.make_store('Group Store B', 'gm@group.test', '22222')
        gm = self.sign_in('gm@group.test')
        # on two stores with none picked yet: asked, never guessed
        for path in ('/reconcile', '/', '/history'):
            r = gm.get(path)
            self.assertEqual(r.status_code, 302, path)
            self.assertIn('/stores', r.headers['Location'])
        page = gm.get('/stores')
        self.assertIn(b'Group Store A', page.data)
        self.assertIn(b'Group Store B', page.data)
        self.assertIn(b'owner@example.com', page.data)      # who added them
        r = self.start_recon(gm, {'dc_text': 'x'})
        self.assertEqual(r.status_code, 302)                # no store, no run
        self.post(gm, '/switch-store', {'store_id': str(b['id'])})
        page = gm.get('/reconcile')
        self.assertEqual(page.status_code, 200)
        with gm.session_transaction() as s:
            self.assertEqual(s['sid'], b['id'])
        # a store the person is not on is ignored
        other = self.make_store('Not Mine', 'x@notmine.test')
        self.post(gm, '/switch-store', {'store_id': str(other['id'])})
        with gm.session_transaction() as s:
            self.assertEqual(s['sid'], b['id'])
        # the next sign-in lands in the store worked in last
        again = self.sign_in('gm@group.test')
        self.assertEqual(again.get('/reconcile').status_code, 200)
        with again.session_transaction() as s:
            self.assertEqual(s['sid'], b['id'])
        self.assertNotEqual(a['id'], b['id'])

    def test_being_added_elsewhere_never_moves_where_work_lands(self):
        from accounts import core
        mine = self.make_store('Victim Store', 'vic@victim.test', '11111')
        self.make_store('Rival Store', 'rival@rival.test', '22222')
        vic = self.sign_in('vic@victim.test')
        self.assertEqual(vic.get('/reconcile').status_code, 200)
        # the rival's admin adds the victim's email to the rival store
        rival = self.sign_in('rival@rival.test')
        self.post(rival, '/users/add', {'email': 'vic@victim.test',
                                        'role': 'user'})
        invite = self.last_link('vic@victim.test')
        # an already signed-in browser keeps working in its own store
        with vic.session_transaction() as s:
            self.assertEqual(s['sid'], mine['id'])
        # a fresh sign-in still lands in the victim's own store
        fresh = self.sign_in('vic@victim.test')
        self.assertEqual(fresh.get('/reconcile').status_code, 200)
        with fresh.session_transaction() as s:
            self.assertEqual(s['sid'], mine['id'])
        # even the rival's invitation link does not drop them in the
        # rival store: they are shown both stores and who added them
        self.post(rival, '/users/add', {'email': 'vic@victim.test',
                                        'role': 'user'})   # no new mail
        clicked = self.client()
        self.assertEqual(clicked.get(invite).status_code, 410)  # retired
        vic_user = core.user_by_email('vic@victim.test')
        self.post(rival, f'/users/{vic_user["id"]}/invite')
        invite = self.last_link('vic@victim.test')
        r = self.use_link(clicked, invite)
        self.assertIn('/stores', r.headers['Location'])
        page = clicked.get('/stores').data
        self.assertIn(b'rival@rival.test', page)
        # the rival's Users page shows only activity in the rival store
        users_page = rival.get('/users').get_data(as_text=True)
        self.assertIn('vic@victim.test', users_page)
        self.assertIn('Not yet', users_page)
        # removing the victim from the rival store does not touch their
        # own store, session or links
        self.post(fresh, '/logout')
        self.post(self.client(), '/login', {'email': 'vic@victim.test'})
        own_link = self.last_link('vic@victim.test')
        self.post(rival, f'/users/{vic_user["id"]}/remove')
        self.assertEqual(self.client().get(own_link).status_code, 200)
        self.assertEqual(vic.get('/reconcile').status_code, 200)

    def test_a_stale_tab_cannot_act_on_another_store(self):
        from accounts import core, storage
        a = self.make_store('Tab Store A', 'al@taba.test', '11111')
        b = self.make_store('Tab Store B', 'bo@tabb.test', '22222')
        owner = self.sign_in(OWNER)
        self.post(owner, f'/owner/open/{a["id"]}')
        token, page_store = self.guard(owner)       # tab 1 shows store A
        self.assertEqual(page_store, str(a['id']))
        self.post(owner, f'/owner/open/{b["id"]}')  # tab 2 opens store B
        stale = {'csrf_token': token, 'store_check': page_store}
        al = core.user_by_email('al@taba.test')
        # every click from the stale tab is refused, nothing changes
        r = owner.post('/users/add', data=dict(
            stale, email='al@taba.test', role='admin'))
        self.assertEqual(r.status_code, 409)
        self.assertIsNone(core.role_in(al['id'], b['id']))
        r = owner.post('/store/settings', data=dict(
            stale, name='Tab Store A', dealer_code='11111',
            threshold_days='60'))
        self.assertEqual(r.status_code, 409)
        self.assertEqual(core.store_by_id(b['id'])['name'], 'Tab Store B')
        r = owner.post('/api/reconcile', data={'dc_text': 'x'},
                       content_type='multipart/form-data',
                       headers={'X-CSRF-Token': token,
                                'X-Store-Id': page_store})
        self.assertEqual(r.status_code, 409)
        self.assertIsNone(storage.current_master(b['id']))
        # a post that names no store at all is refused too
        r = owner.post('/users/add', data={
            'csrf_token': token, 'email': 'al@taba.test', 'role': 'admin'})
        self.assertEqual(r.status_code, 409)
        # after a reload the page matches and the click goes through
        r = self.post(owner, '/users/add', {'email': 'new@tabb.test',
                                            'role': 'user'})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(core.role_in(
            core.user_by_email('new@tabb.test')['id'], b['id']), 'user')


class TestSavedWork(AccountsCase):
    def test_run_is_saved_and_next_run_starts_from_the_saved_master(self):
        from accounts import db, storage
        from recon import jobs
        store = self.make_store('Saved Store', 'sue@saved.test')
        sue = self.sign_in('sue@saved.test')
        self.assertIsNone(storage.current_master(store['id']))

        job1, job = self.first_run(sue)
        self.assertEqual(job['status'], 'done', job.get('error'))
        self.assertTrue(job['result']['saved'])
        self.assertEqual(job['result']['master_held'], '')
        self.assertNotIn('store_id', job)        # ownership stays server-side
        m1 = storage.current_master(store['id'])
        self.assertEqual((m1['row_count'], m1['unpaid_count']), (1, 1))
        self.assertEqual(float(m1['unpaid_amount']), 50.0)
        self.assertEqual(m1['created_email'], 'sue@saved.test')
        self.assertEqual(m1['source'], 'run')

        # second run: no workbook uploaded, so it starts from the saved one
        page = sue.get('/reconcile')
        self.assertIn(b"Saved on your store", page.data)
        job2, job = self.run_recon(sue, {
            'dc_text': 'C400000001\t11111111AA\t\t50.00\t\t\t',
            'asof': '2026-10-06', 'threshold': '60'})
        self.assertEqual(job['status'], 'done', job.get('error'))
        self.assertEqual(job['result']['inputs']['master_rows'], 1)
        run2 = storage.get_run(job2, store['id'])
        self.assertEqual(run2['master_in'], m1['id'])
        m2 = storage.current_master(store['id'])
        self.assertEqual(run2['master_out'], m2['id'])
        self.assertNotEqual(m1['id'], m2['id'])
        self.assertEqual(len(storage.master_versions(store['id'])), 2)
        self.assertEqual(run2['user_email'], 'sue@saved.test')
        self.assertEqual(run2['summary']['credit_request'], 0)  # asked 10/05

        # history lists both, and survives a server restart
        hist = sue.get('/history')
        self.assertEqual(hist.status_code, 200)
        self.assertEqual(hist.data.count(b'/reconcile#job='), 2)
        with jobs._lock:
            jobs._jobs.clear()
            jobs._meta.clear()
        shutil.rmtree(jobs.job_dir(job1), ignore_errors=True)
        again = sue.get(f'/api/jobs/{job1}').get_json()
        self.assertEqual(again['status'], 'done')
        self.assertEqual([x['ticket'] for x in
                          again['result']['credit_request']], ['C400000001'])
        names = {f['name']: f for f in again['result']['files']}
        dl = sue.get(f'/api/jobs/{job1}/files/Credit_Request.xlsx')
        self.assertEqual(dl.status_code, 200)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(dl.data))['Credit Request']
        self.assertEqual([c.value for c in ws[2]],
                         [12345, '2026-07-01', 'C400000001'])
        zname = next(n for n, f in names.items() if f['kind'] == 'zip')
        z = sue.get(f'/api/jobs/{job1}/files/{zname}')
        self.assertEqual(z.status_code, 200)
        import zipfile
        inside = set(zipfile.ZipFile(io.BytesIO(z.data)).namelist())
        self.assertIn('Credit_Request.xlsx', inside)
        self.assertIn('Core_Returns_RECONCILED.xlsx', inside)
        self.assertEqual(sue.get(
            f'/api/jobs/{job1}/files/../app.py').status_code, 404)
        self.assertEqual(sue.get(
            f'/api/jobs/{job1}/files/nope.xlsx').status_code, 404)
        self.assertEqual(sue.get('/master/download').status_code, 200)
        self.assertEqual(
            db.one('SELECT count(*) AS n FROM runs WHERE store_id = %s',
                   (store['id'],))['n'], 2)

    def test_one_store_never_sees_another_stores_work(self):
        from accounts import storage
        a = self.make_store('Iso Store A', 'al@isoa.test')
        b = self.make_store('Iso Store B', 'bea@isob.test')
        al = self.sign_in('al@isoa.test')
        bea = self.sign_in('bea@isob.test')
        job_a, job = self.first_run(al, 'C400000011')
        self.assertEqual(job['status'], 'done', job.get('error'))
        job_b, job = self.first_run(bea, 'C400000022')
        self.assertEqual(job['status'], 'done', job.get('error'))
        ma = storage.current_master(a['id'])
        mb = storage.current_master(b['id'])

        # live (still in the server's memory)
        self.assertEqual(bea.get(f'/api/jobs/{job_a}').status_code, 404)
        self.assertEqual(bea.get(
            f'/api/jobs/{job_a}/files/Credit_Request.xlsx').status_code, 404)
        # from history
        from recon import jobs
        with jobs._lock:
            jobs._jobs.clear()
            jobs._meta.clear()
        self.assertEqual(bea.get(f'/api/jobs/{job_a}').status_code, 404)
        self.assertEqual(bea.get(
            f'/api/jobs/{job_a}/files/Credit_Request.xlsx').status_code, 404)
        self.assertEqual(bea.get(f'/api/jobs/{job_b}').status_code, 200)
        # masters
        self.assertEqual(bea.get(
            f'/master/{ma["id"]}/download').status_code, 404)
        self.assertEqual(self.post(
            bea, f'/master/{ma["id"]}/restore').status_code, 404)
        self.assertEqual(bea.get(
            f'/master/{mb["id"]}/download').status_code, 200)
        # history and the current master are the store's own
        hist = bea.get('/history').data
        self.assertIn(job_b.encode(), hist)
        self.assertNotIn(job_a.encode(), hist)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(
            bea.get('/master/download').data))['Core Returns']
        self.assertEqual(ws['C2'].value, 'C400000022')
        # the owner sees a store only after opening it
        owner = self.sign_in(OWNER)
        self.post(owner, f'/owner/open/{a["id"]}')
        self.assertEqual(owner.get(f'/api/jobs/{job_a}').status_code, 200)
        self.assertEqual(owner.get(f'/api/jobs/{job_b}').status_code, 404)

    def test_a_failed_run_leaves_the_master_alone(self):
        from accounts import storage
        store = self.make_store('Fail Store', 'fay@fail.test')
        fay = self.sign_in('fay@fail.test')
        self.first_run(fay)
        before = storage.current_master(store['id'])
        job_id, job = self.run_recon(fay, {
            'master': (io.BytesIO(b'not a workbook'), 'broken.xlsx'),
            'asof': '2026-10-06'})
        self.assertEqual(job['status'], 'error')
        self.assertEqual(storage.current_master(store['id'])['id'],
                         before['id'])
        run = storage.get_run(job_id, store['id'])
        self.assertEqual(run['status'], 'error')
        self.assertIn(b'Stopped', fay.get('/history').data)

    def test_only_a_real_master_from_an_admin_replaces_the_saved_one(self):
        from accounts import storage
        from openpyxl import Workbook
        store = self.make_store('Swap Store', 'sal@swap.test')
        sal = self.sign_in('sal@swap.test')
        self.first_run(sal)
        before = storage.current_master(store['id'])
        self.post(sal, '/users/add', {'email': 'tom@swap.test',
                                      'role': 'user'})
        tom = self.sign_in('tom@swap.test')
        other = master_bytes([row(None, '2026-07-02 17:00.01', 'C400000077',
                                  '22222222AA', 75)])
        # a plain user cannot swap the saved master for another workbook
        self.assertNotIn(b'Use a different workbook',
                         tom.get('/reconcile').data)
        r = self.start_recon(tom, {
            'master': (io.BytesIO(other), 'Core_Returns_RECONCILED.xlsx'),
            'asof': '2026-10-06'})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(storage.current_master(store['id'])['id'],
                         before['id'])
        # but runs fine from the saved one
        _, job = self.run_recon(tom, {
            'dc_text': 'C400000001\t11111111AA\t\t50.00\t\t\t',
            'asof': '2026-10-06'})
        self.assertEqual(job['status'], 'done', job.get('error'))
        middle = storage.current_master(store['id'])
        # an admin picking one of the RESULT workbooks by mistake: refused
        wb = Workbook()
        ws = wb.active
        ws.title = 'Uncredited'
        ws.append(['Shipment ID', 'Control Ticket', 'Core Part #'])
        buf = io.BytesIO()
        wb.save(buf)
        _, job = self.run_recon(sal, {
            'master': (io.BytesIO(buf.getvalue()), 'Core_Recon_Report.xlsx'),
            'asof': '2026-10-06'})
        self.assertEqual(job['status'], 'error')
        self.assertIn('not a Core Returns master', job['error'])
        self.assertEqual(storage.current_master(store['id'])['id'],
                         middle['id'])
        # an admin with a real master: replaced
        _, job = self.run_recon(sal, {
            'master': (io.BytesIO(other), 'Core_Returns_RECONCILED.xlsx'),
            'asof': '2026-10-06'})
        self.assertEqual(job['status'], 'done', job.get('error'))
        self.assertNotEqual(storage.current_master(store['id'])['id'],
                            middle['id'])

    def test_first_master_can_be_the_stores_own_workbook_layout(self):
        from accounts import storage
        from openpyxl import Workbook
        store = self.make_store('Layout Store', 'lee@layout.test')
        lee = self.sign_in('lee@layout.test')
        wb = Workbook()
        ws = wb.active                       # tab is "Sheet", not Core Returns
        ws.append(['Credit Memo #', 'Shipment ID', 'Control Ticket',
                   'Core Part #', 'Claim #', 'Credit Amount'])
        ws.append([None, '2026-09-01 17:00.01', 'C400000051', '11111111AA',
                   None, 50])
        buf = io.BytesIO()
        wb.save(buf)
        _, job = self.run_recon(lee, {
            'master': (io.BytesIO(buf.getvalue()), 'our cores.xlsx'),
            'asof': '2026-10-06'})
        self.assertEqual(job['status'], 'done', job.get('error'))
        self.assertEqual(storage.current_master(store['id'])['row_count'], 1)

    def test_a_much_smaller_master_is_held_for_an_admin_to_confirm(self):
        from accounts import storage
        store = self.make_store('Shrink Store', 'shay@shrink.test')
        shay = self.sign_in('shay@shrink.test')
        _, job = self.run_recon(shay, {
            'master': (io.BytesIO(master_bytes(many_rows(20))),
                       'Core_Returns_RECONCILED.xlsx'),
            'asof': '2026-10-06'})
        self.assertEqual(job['status'], 'done', job.get('error'))
        big = storage.current_master(store['id'])
        self.assertEqual(big['row_count'], 20)
        # the wrong (tiny) workbook
        job_id, job = self.run_recon(shay, {
            'master': (io.BytesIO(master_bytes(many_rows(2, 500))),
                       'Core_Returns_RECONCILED.xlsx'),
            'asof': '2026-10-06'})
        self.assertEqual(job['status'], 'done', job.get('error'))
        self.assertTrue(job['result']['saved'])
        self.assertIn('far fewer rows', job['result']['master_held'])
        self.assertEqual(storage.current_master(store['id'])['id'],
                         big['id'])                      # untouched
        held = storage.master_versions(store['id'])[0]
        self.assertFalse(held['is_current'])
        self.assertIn('far fewer rows', held['held_reason'])
        page = shay.get('/history').data
        self.assertIn(b'Not made current', page)
        self.assertIn(b'Master held', page)
        # the warning is still there when the run is opened from History
        from recon import jobs
        with jobs._lock:
            jobs._jobs.clear()
            jobs._meta.clear()
        again = shay.get(f'/api/jobs/{job_id}').get_json()['result']
        self.assertIn('far fewer rows', again['master_held'])
        self.assertTrue(again['from_history'])
        # the admin decides it was right after all
        self.assertEqual(self.post(
            shay, f'/master/{held["id"]}/restore').status_code, 302)
        self.assertEqual(storage.current_master(store['id'])['row_count'],
                         2)

    def test_a_run_never_overwrites_a_master_changed_under_it(self):
        from accounts import storage
        from recon import jobs
        store = self.make_store('Race Store', 'ray@race.test')
        ray = self.sign_in('ray@race.test')
        job1, _ = self.first_run(ray, 'C400000041')
        v1 = storage.current_master(store['id'])
        self.run_recon(ray, {
            'dc_text': 'C400000041\t11111111AA\t\t50.00\t\t\t',
            'asof': '2026-10-06'})
        v2 = storage.current_master(store['id'])
        # a run started from v2 is still working when an admin puts v1
        # back; when the run finishes it must not undo that
        actor = {'id': None, 'email': 'ray@race.test'}
        storage.begin_run('b' * 32, store['id'], actor, 'reconcile')
        v3 = storage.restore_master(store['id'], v1['id'], actor)
        out = os.path.join(jobs.job_dir(job1), 'out')
        result = ray.get(f'/api/jobs/{job1}').get_json()['result']
        saved = storage.finish_run('b' * 32, store['id'], actor, result,
                                   out, master_in=v2['id'])
        self.assertIn('changed by someone else', saved['held'])
        self.assertEqual(storage.current_master(store['id'])['id'], v3)
        kept = storage.master_file(store['id'], saved['master_out'])[0]
        self.assertFalse(kept['is_current'])
        # a plain user's uploaded workbook that finishes after a master
        # was saved (the first-run race) is held as well
        storage.begin_run('c' * 32, store['id'], actor, 'reconcile')
        saved = storage.finish_run('c' * 32, store['id'], actor, result,
                                   out, master_in=None, may_replace=False)
        self.assertIn('only a store admin', saved['held'])
        self.assertEqual(storage.current_master(store['id'])['id'], v3)

    def test_an_admin_can_put_an_earlier_master_back(self):
        from accounts import storage
        store = self.make_store('Undo Store', 'una@undo.test')
        una = self.sign_in('una@undo.test')
        self.first_run(una, 'C400000031')
        v1 = storage.current_master(store['id'])
        # a run with the wrong workbook replaces the master...
        wrong = master_bytes([row(None, '2026-07-02 17:00.01', 'C400000099',
                                  '22222222AA', 75)])
        self.run_recon(una, {
            'master': (io.BytesIO(wrong), 'Core_Returns_RECONCILED.xlsx'),
            'asof': '2026-10-06'})
        v2 = storage.current_master(store['id'])
        self.assertNotEqual(v1['id'], v2['id'])
        # ...a plain user cannot undo it, an admin can
        self.post(una, '/users/add', {'email': 'uri@undo.test',
                                      'role': 'user'})
        uri = self.sign_in('uri@undo.test')
        self.assertEqual(self.post(
            uri, f'/master/{v1["id"]}/restore').status_code, 403)
        self.assertEqual(self.post(
            una, f'/master/{v1["id"]}/restore').status_code, 302)
        v3 = storage.current_master(store['id'])
        self.assertEqual(v3['source'], 'restore')
        _, now = storage.master_file(store['id'])
        _, first = storage.master_file(store['id'], v1['id'])
        self.assertEqual(now, first)
        self.assertEqual(len(storage.master_versions(store['id'])), 3)

    def test_a_run_that_cannot_be_saved_says_so(self):
        from accounts import storage
        store = self.make_store('Unsaved Store', 'uma@unsaved.test')
        uma = self.sign_in('uma@unsaved.test')
        real = storage.finish_run

        def boom(*a, **k):
            raise RuntimeError('database went away')

        storage.finish_run = boom
        try:
            job_id, job = self.first_run(uma)
        finally:
            storage.finish_run = real
        self.assertEqual(job['status'], 'done')
        self.assertFalse(job['result']['saved'])
        self.assertIn('could not be saved', job['result']['warnings'][0])
        self.assertIsNone(storage.current_master(store['id']))
        self.assertEqual(storage.get_run(job_id, store['id'])['status'],
                         'error')                    # not stuck on Running

    def test_quick_check_is_logged(self):
        from accounts import storage
        store = self.make_store('Quick Store', 'quin@quick.test')
        storage.log_quick_check('a' * 32, store['id'],
                                {'id': None, 'email': 'quin@quick.test'},
                                {'shipper_count': 25, 'credited_count': 4,
                                 'unclaimed_count': 21,
                                 'unclaimed_total': 4740.0})
        quin = self.sign_in('quin@quick.test')
        page = quin.get('/history').data
        self.assertIn(b'Quick Check', page)
        self.assertIn(b'$4,740.00', page)
        # a quick check is a history line only: nothing to open
        self.assertEqual(quin.get('/api/jobs/' + 'a' * 32).status_code, 404)


if __name__ == '__main__':
    unittest.main()
