"""What a store keeps between visits: its master workbook (every version)
and the history of its runs with their result files.

Files are stored in the database (they are small: workbooks and a few PDF
pages per run), so the web service itself stays stateless.
"""
import io
import os
import zipfile

from . import db
from .core import _audit

# Largest single result file kept in history (anything bigger is skipped
# and stays downloadable only until the server's 2-hour cleanup).
MAX_FILE_BYTES = 40 * 1024 * 1024
# A file is stored in pieces of this size, one piece per statement. The
# database needs about four times a statement's size in memory while it
# takes it in, and the smallest database plan has 256 MB in all: a 30 MB
# scan-heavy PDF in one statement could knock it over.
PART_BYTES = 2 * 1024 * 1024
# Result files are kept for this many most-recent finished runs per store.
# Older runs stay in the history list (who, when, the numbers) without
# files. Masters are never pruned.
KEEP_RUN_FILES = 104
# A run whose master comes out with fewer than this share of the saved
# master's rows is not made current on its own (wrong workbook picked).
SHRINK_GUARD = 0.5
SHRINK_MIN_ROWS = 10

INTERRUPTED = ('The server restarted while this run was working. Nothing '
               'was saved. Start the run again.')
HELD_CHANGED = ('the saved master was changed by someone else while this '
                'run was waiting or working')
STOPPED_UNSAVED = ('This run stopped before it could be saved. Start it '
                   'again.')
HELD_NOT_ADMIN = ('a master was saved for this store while the run was '
                  'working, and only a store admin can replace a saved '
                  'master')
HELD_SHRANK = ('it has far fewer rows than the saved master ({new:,} '
               'against {old:,}), which usually means the wrong workbook '
               'was used')


class NotAMaster(Exception):
    pass


# ── masters ─────────────────────────────────────────────────────────────
_MASTER_COLS = ('id, store_id, filename, row_count, unpaid_count, '
                'unpaid_amount, source, held_reason, run_id, created_by, '
                'created_email, created_at, is_current, '
                'octet_length(content) AS size')


def current_master(store_id):
    """The store's current master (no file bytes), or None."""
    return db.one(f'SELECT {_MASTER_COLS} FROM masters WHERE store_id = %s '
                  f'AND is_current', (store_id,))


def master_versions(store_id, limit=30):
    return db.rows(f'SELECT {_MASTER_COLS} FROM masters WHERE store_id = %s '
                   f'ORDER BY id DESC LIMIT %s', (store_id, limit))


def master_file(store_id, master_id=None):
    """(row, bytes) of one master version (the current one by default).
    Always filtered by store, so one store can never read another's."""
    if master_id is None:
        row = db.one('SELECT * FROM masters WHERE store_id = %s AND '
                     'is_current', (store_id,))
    else:
        row = db.one('SELECT * FROM masters WHERE store_id = %s AND '
                     'id = %s', (store_id, master_id))
    if not row:
        return None, None
    content = bytes(row.pop('content'))
    return row, content


def check_master_upload(path):
    """Refuse a workbook that is not a Core Returns master before it can
    become a store's saved master. The engine itself is more forgiving
    (any sheet with a Control Ticket column), which is how one of the
    result workbooks could be picked by mistake."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True)
        names = list(wb.sheetnames)
        wb.close()
    except Exception:
        raise NotAMaster(
            'The master workbook could not be opened as an Excel file. '
            'Open it in Excel, choose Save As > Excel Workbook (.xlsx), '
            'and add it again.')
    if 'Core Returns' not in names:
        raise NotAMaster(
            'That workbook is not a Core Returns master: it has no "Core '
            'Returns" sheet. Pick Core_Returns_RECONCILED.xlsx (not one of '
            'the result workbooks).')


def _master_stats(content):
    """(rows, unpaid_count, unpaid_amount) read from a master workbook."""
    import tempfile
    from recon import master as M
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    try:
        with os.fdopen(fd, 'wb') as fh:
            fh.write(content)
        rows, _ = M.read_master(path)
    finally:
        try:
            os.remove(path)
        except OSError:
            pass
    unpaid = [r for r in rows if not r.get('memo')]
    amount = 0.0
    for r in unpaid:
        try:
            amount += float(r.get('amount') or 0)
        except (TypeError, ValueError):
            pass
    return len(rows), len(unpaid), round(amount, 2)


def _lock_masters(conn, store_id):
    """One change to a store's masters at a time, held to the end of the
    transaction. Returns the current master's (id, row_count) or None."""
    conn.execute('SELECT pg_advisory_xact_lock(712005, %s::int)',
                 (int(store_id) % 2147483647,))
    return conn.execute('SELECT id, row_count FROM masters WHERE '
                        'store_id = %s AND is_current',
                        (store_id,)).fetchone()


def _insert_master(conn, store_id, filename, content, stats, source, actor,
                   run_id, make_current, held_reason=''):
    n_rows, n_unpaid, amt = stats
    if make_current:
        conn.execute('UPDATE masters SET is_current = FALSE WHERE '
                     'store_id = %s AND is_current', (store_id,))
    return conn.execute(
        'INSERT INTO masters (store_id, filename, content, row_count, '
        'unpaid_count, unpaid_amount, source, held_reason, run_id, '
        'created_by, created_email, is_current) VALUES (%s, %s, %s, %s, '
        '%s, %s, %s, %s, %s, %s, %s, %s) RETURNING id',
        (store_id, filename or 'Core_Returns_RECONCILED.xlsx', content,
         n_rows, n_unpaid, amt, source, held_reason, run_id,
         (actor or {}).get('id'), (actor or {}).get('email', ''),
         make_current)).fetchone()['id']


def save_master(store_id, filename, content, source, actor, run_id=None):
    """Save a new version and make it the current one. Returns its id."""
    stats = _master_stats(content)
    with db.connect() as conn:
        _lock_masters(conn, store_id)
        return _insert_master(conn, store_id, filename, content, stats,
                              source, actor, run_id, True)


def restore_master(store_id, master_id, actor):
    """Make an older version current again (as a new version, so nothing
    is ever lost and the restore itself can be undone)."""
    row, content = master_file(store_id, master_id)
    if not row:
        return None
    stats = _master_stats(content)
    with db.connect() as conn:
        _lock_masters(conn, store_id)
        new_id = _insert_master(conn, store_id, row['filename'], content,
                                stats, 'restore', actor, row.get('run_id'),
                                True)
        _audit(conn, store_id, actor, 'master_restored',
               {'from_version': master_id, 'new_version': new_id})
    return new_id


# ── runs ────────────────────────────────────────────────────────────────
def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def summarize(result):
    """The handful of numbers shown on the history list."""
    r = result or {}

    def total(key, field):
        return round(sum(_num(x.get(field)) for x in r.get(key, [])), 2)

    return {
        'matched': len(r.get('matched', [])),
        'matched_amount': total('matched', 'credited'),
        'new_shipments': len(r.get('new_shipments', [])),
        'new_amount': total('new_shipments', 'amount'),
        'unpaid': len(r.get('unpaid', [])),
        'unpaid_amount': total('unpaid', 'amount'),
        'credit_request': len(r.get('credit_request', [])),
        'request_amount': total('credit_request', 'amount'),
        'pages_scanned': r.get('pages_scanned', 0),
        'warnings': len(r.get('warnings', [])),
        'dealer': r.get('dealer') or '',
        'memos': [m.get('memo') for m in r.get('memo_totals', [])][:20],
    }


def begin_run(run_id, store_id, actor, kind, asof=None, options=None):
    db.execute(
        'INSERT INTO runs (id, store_id, user_id, user_email, kind, status, '
        'asof, options) VALUES (%s, %s, %s, %s, %s, %s, %s, %s) '
        'ON CONFLICT (id) DO NOTHING',
        (run_id, store_id, (actor or {}).get('id'),
         (actor or {}).get('email', ''), kind, 'running', asof,
         db.jsonb(options or {})))


def close_interrupted_runs():
    db.execute("UPDATE runs SET status = 'error', error = %s, "
               "finished_at = now() WHERE status = 'running'",
               (INTERRUPTED,))


def fail_run(run_id, error):
    db.execute("UPDATE runs SET status = 'error', error = %s, "
               "finished_at = now() WHERE id = %s", (str(error)[:2000],
                                                      run_id))


def finish_run(run_id, store_id, actor, result, out_dir, master_in=None,
               may_replace=True, seen_master=None):
    """Store a finished run: its numbers, the full result, the files, and
    the new master. All in one transaction.

    master_in is the saved master the run started from, or None when it
    started from an uploaded workbook (or from nothing). seen_master is
    the master that was current when the person pressed Run.

    The new master becomes the store's current one unless something says
    it should not on its own:
    - the run started from the saved master and that master was changed
      (restored, replaced) while the run was working, or
    - the run started from an uploaded workbook and the person is not an
      admin (may_replace False) while the store has a saved master by
      now, or
    - the run started from an uploaded workbook and a different master
      is current than the one the person was looking at when they pressed
      Run (another run saved while this one waited its turn), or
    - it came out with far fewer rows than the saved master.
    Then it is kept as a version an admin can make current from History.

    Returns {'master_out': id or None, 'held': reason or ''}.
    """
    files = []
    master_name = None
    too_big = []
    for f in (result or {}).get('files', []):
        if f.get('kind') == 'zip':
            continue                    # rebuilt from the others on demand
        path = os.path.join(out_dir, f['name'])
        if not os.path.isfile(path):
            continue
        size = os.path.getsize(path)
        if size > MAX_FILE_BYTES:
            too_big.append(f['name'])
            continue
        with open(path, 'rb') as fh:
            content = fh.read()
        files.append((f['name'], f.get('label', ''), f.get('kind', ''),
                      size, content))
        if f.get('kind') == 'master':
            master_name = f['name']
    master_content = next((c for n, _, _, _, c in files
                           if n == master_name), None)
    stats = _master_stats(master_content) if master_content else None
    held = ''
    with db.connect() as conn:
        master_out = None
        if master_content is not None:
            current = _lock_masters(conn, store_id)
            if current:
                if master_in is not None and current['id'] != master_in:
                    held = HELD_CHANGED
                elif master_in is None and not may_replace:
                    held = HELD_NOT_ADMIN
                elif master_in is None and current['id'] != seen_master:
                    held = HELD_CHANGED
                elif current['row_count'] >= SHRINK_MIN_ROWS and \
                        stats[0] < current['row_count'] * SHRINK_GUARD:
                    held = HELD_SHRANK.format(new=stats[0],
                                              old=current['row_count'])
            master_out = _insert_master(
                conn, store_id, master_name, master_content, stats, 'run',
                actor, run_id, make_current=not held, held_reason=held)
        for name, label, kind, size, content in files:
            added = conn.execute(
                'INSERT INTO run_files (run_id, name, label, kind, size, '
                "content) VALUES (%s, %s, %s, %s, %s, ''::bytea) "
                'ON CONFLICT (run_id, name) DO NOTHING',
                (run_id, name, label, kind, size)).rowcount
            if not added:
                continue
            for n, at in enumerate(range(0, len(content), PART_BYTES)):
                conn.execute(
                    'INSERT INTO run_file_parts (run_id, name, part, '
                    'content) VALUES (%s, %s, %s, %s)',
                    (run_id, name, n, content[at:at + PART_BYTES]))
        kept = dict(result or {}, saved=True, master_held=held)
        if too_big:
            # said on the page when the run is opened from History
            kept['files'] = [dict(f, kept=f['name'] not in too_big)
                             for f in kept.get('files', [])]
        conn.execute(
            "UPDATE runs SET status = 'done', summary = %s, result = %s, "
            "master_in = %s, master_out = %s, finished_at = now() "
            "WHERE id = %s",
            (db.jsonb(dict(summarize(result), master_held=held)),
             db.jsonb(kept), master_in, master_out, run_id))
        # keep files only for this store's most recent finished runs
        conn.execute(
            "DELETE FROM run_files WHERE run_id IN (SELECT id FROM runs "
            "WHERE store_id = %s AND kind = 'reconcile' AND "
            "status = 'done' ORDER BY created_at DESC OFFSET %s)",
            (store_id, KEEP_RUN_FILES))
    return {'master_out': master_out, 'held': held}


def log_quick_check(run_id, store_id, actor, summary):
    db.execute(
        "INSERT INTO runs (id, store_id, user_id, user_email, kind, status, "
        "summary, finished_at) VALUES (%s, %s, %s, %s, 'quick', 'done', "
        "%s, now()) ON CONFLICT (id) DO NOTHING",
        (run_id, store_id, (actor or {}).get('id'),
         (actor or {}).get('email', ''), db.jsonb(summary)))


def list_runs(store_id, limit=100):
    return db.rows(
        'SELECT id, user_email, kind, status, asof, summary, error, '
        'created_at, finished_at, master_in, master_out, '
        '(SELECT count(*) FROM run_files f WHERE f.run_id = runs.id) '
        'AS file_count FROM runs WHERE store_id = %s '
        'ORDER BY created_at DESC LIMIT %s', (store_id, limit))


def get_run(run_id, store_id):
    """One run of one store. Never without the store."""
    return db.one('SELECT * FROM runs WHERE id = %s AND store_id = %s',
                  (run_id, store_id))


def run_file(run_id, name):
    """bytes of one stored result file, or None."""
    row = db.one('SELECT content, size FROM run_files WHERE run_id = %s '
                 'AND name = %s', (run_id, name))
    if not row:
        return None
    whole = bytes(row['content'])
    if whole or not row['size']:
        return whole            # kept in one piece (or an empty file)
    parts = db.rows('SELECT content FROM run_file_parts WHERE run_id = %s '
                    'AND name = %s ORDER BY part', (run_id, name))
    return b''.join(bytes(p['content']) for p in parts)


def run_zip(run_id):
    """The run's "everything" zip, rebuilt from its stored files."""
    names = [r['name'] for r in db.rows(
        'SELECT name FROM run_files WHERE run_id = %s ORDER BY name',
        (run_id,))]
    if not names:
        return None
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for name in names:
            z.writestr(name, run_file(run_id, name) or b'')
    return buf.getvalue()
