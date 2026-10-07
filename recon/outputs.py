"""Output files: Credit Request (xlsx + marked shipper PDF), Shippers To
Re-Scan (xlsx + pdf), DealerCONNECT comparison and the run report."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = 'FF1F4E78'
THIN = Side(style='thin', color='FFB7C3D0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFONT = Font(name='Arial', size=10, bold=True, color='FFFFFFFF')
HFILL = PatternFill('solid', fgColor=NAVY)
DFONT = Font(name='Arial', size=10)
BFONT = Font(name='Arial', size=10, bold=True)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=False)
WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True)
MONEY = '"$"#,##0.00'
STATUS_FILL = {
    'ADD TO MASTER': 'FFFFE699', 'MOPAR HAS NO RECORD': 'FFF8B7B7',
    'NOT ON MOPAR UNPAID LIST': 'FFF8CBAD', 'FIX DATA': 'FFDDEBF7',
    'FIXED FROM DEALERCONNECT': 'FFDDEBF7', 'PAID PER MOPAR': 'FFC6EFCE',
    'CHECK - CREDITED ON MASTER': 'FFF8CBAD', 'AGREE': None,
}


def table_sheet(ws, headers, rows, widths=None, money_cols=(), center=True,
                status_col=None):
    """Write a styled table (header row 1, freeze, filter)."""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = \
            HFONT, HFILL, CENTER, BORDER
    ws.row_dimensions[1].height = 22
    for r, vals in enumerate(rows, start=2):
        fill = None
        if status_col is not None:
            color = STATUS_FILL.get(vals[status_col])
            fill = PatternFill('solid', fgColor=color) if color else None
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = DFONT, BORDER
            cell.alignment = CENTER if center else LEFT
            if c - 1 in money_cols:
                cell.number_format = MONEY
            if fill:
                cell.fill = fill
    for c, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = (f'A1:{get_column_letter(len(headers))}'
                              f'{len(rows) + 1}')


def write_credit_request_xlsx(request, dealer, path):
    """Same 3 columns as the Credit_Request.xlsx sent to Mopar."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Credit Request'
    dc = int(dealer) if dealer and str(dealer).isdigit() else (dealer or '')
    rows = [[dc, r['date'], r['ticket']] for r in request]
    table_sheet(ws, ['Dealer Code', 'GCRS Shipper Date',
                     'Core Return Material Ticket #'], rows, [14, 20, 30])
    for r in range(2, len(rows) + 2):
        for c in range(1, 4):
            ws.cell(row=r, column=c).fill = PatternFill('solid',
                                                        fgColor='FFFFFFFF')
    wb.save(path)


def write_rescan_xlsx(rescan, note, path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Shippers To Re-Scan'
    if rescan:
        rows = [[g['date'], ', '.join(g['tickets'])] for g in rescan]
    else:
        rows = [['(none)', note]]
    table_sheet(ws, ['GCRS Shipper Date', 'Control Ticket Numbers'], rows,
                [20, 90], center=False)
    wb.save(path)


def write_rescan_pdf(rescan, note, asof, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=landscape(letter),
                            leftMargin=40, rightMargin=40, topMargin=40,
                            bottomMargin=40)
    story = [Paragraph('Shippers To Re-Scan', styles['Title']),
             Paragraph(f'As of {asof}', styles['Normal']), Spacer(1, 12)]
    if not rescan:
        story.append(Paragraph(note, styles['Normal']))
    else:
        story.append(Paragraph(note, styles['Normal']))
        story.append(Spacer(1, 10))
        data = [['GCRS Shipper Date', 'Control Ticket Numbers']]
        for g in rescan:
            data.append([g['date'], Paragraph(', '.join(g['tickets']),
                                              styles['Normal'])])
        t = Table(data, colWidths=[130, 560], repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#B7C3D0')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(t)
    doc.build(story)


def write_credit_request_pdf(pages, path):
    """pages: [{'path', 'page', 'width', 'height', 'boxes': [(l,t,r,b)],
    'checkoff_x'}] in output order. Copies each signed shipper page as-is
    and stamps a red "UnPaid" in the PART CHECKOFF column next to every
    requested ticket on it."""
    from pypdf import PdfReader, PdfWriter, Transformation
    from reportlab.pdfgen import canvas

    writer = PdfWriter()
    readers = {}
    for spec in pages:
        reader = readers.get(spec['path'])
        if reader is None:
            reader = readers[spec['path']] = PdfReader(spec['path'])
        if spec['page'] - 1 >= len(reader.pages):
            continue
        writer.add_page(reader.pages[spec['page'] - 1])
        page = writer.pages[-1]
        if page.get('/Rotate'):
            page.transfer_rotation_to_content()
        mb = page.mediabox
        pw, ph = float(mb.width), float(mb.height)
        sx = pw / float(spec['width'] or 1)
        sy = ph / float(spec['height'] or 1)
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(pw, ph))
        c.setFillColorRGB(0.85, 0.0, 0.0)
        for (l, t, r, b) in spec['boxes']:
            h_pt = max(8.0, (b - t) * sy * 1.25)
            c.setFont('Helvetica-Bold', h_pt)
            text_w = c.stringWidth('UnPaid', 'Helvetica-Bold', h_pt)
            right = l * sx - 0.08 * 72          # just left of the ticket
            x = right - text_w
            if spec.get('checkoff_x') is not None:
                x = max(x, spec['checkoff_x'] * sx - 4)
            x = max(4.0, x)
            y = ph - b * sy + (b - t) * sy * 0.12
            c.drawString(x, y, 'UnPaid')
        c.save()
        buf.seek(0)
        overlay = PdfReader(buf).pages[0]
        page.merge_transformed_page(
            overlay, Transformation().translate(float(mb.left),
                                                float(mb.bottom)))
    with open(path, 'wb') as fh:
        writer.write(fh)


def write_comparison_xlsx(comp, summary, info, path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Action Summary'
    ws['A1'] = 'Core Returns — master vs. DealerCONNECT'
    ws['A1'].font = Font(name='Arial', size=13, bold=True)
    ws['A2'] = info
    ws['A2'].font = DFONT
    ws['A3'] = ('Rule: DealerCONNECT is authoritative for part number, '
                'amount and warranty claim. The master is authoritative for '
                'shipment date.')
    ws['A3'].font = Font(name='Arial', size=10, italic=True)
    hdr = ['Status', 'What it means', 'Tickets', 'Amount', 'Action to take']
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = \
            HFONT, HFILL, CENTER, BORDER
    r = 6
    for s in summary:
        vals = [s['status'], s['meaning'], s['count'], s['amount'],
                s['action']]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = DFONT, BORDER
            cell.alignment = WRAP
            if c == 4:
                cell.number_format = MONEY
        color = STATUS_FILL.get(s['status'])
        if color:
            ws.cell(row=r, column=1).fill = PatternFill('solid',
                                                        fgColor=color)
        r += 1
    ws.cell(row=r, column=1, value='TOTAL').font = BFONT
    ws.cell(row=r, column=3, value=f'=SUM(C6:C{r - 1})').font = BFONT
    tot = ws.cell(row=r, column=4, value=f'=SUM(D6:D{r - 1})')
    tot.font, tot.number_format = BFONT, MONEY
    for col, w in zip('ABCDE', [30, 46, 10, 14, 70]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet('Comparison')
    rows = [[x['status'], x['ticket'], x['date'], x['age'], x['part_ours'],
             x['part_dc'], x['claim'], x['amount_ours'], x['amount_dc'],
             x['requested'], x['note']] for x in comp]
    table_sheet(ws2, ['Status', 'Control Ticket', 'Shipment Date',
                      'Age (days)', 'Part # (ours)', 'Part # (DealerCONNECT)',
                      'Warranty Claim', 'Amount (ours)',
                      'Amount (DealerCONNECT)', 'Credit Requested',
                      'Action / Note'], rows,
                [26, 14, 14, 10, 14, 20, 14, 13, 20, 15, 80],
                money_cols=(7, 8), center=False, status_col=0)
    wb.save(path)


def write_report_xlsx(sheets, path):
    """sheets: list of (title, headers, rows, widths, money_cols)."""
    wb = Workbook()
    first = True
    for title, headers, rows, widths, money_cols in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title[:31]
        if headers is None:      # free-form summary sheet
            for r, vals in enumerate(rows, start=1):
                for c, v in enumerate(vals, start=1):
                    cell = ws.cell(row=r, column=c, value=v)
                    cell.font = DFONT
                    if c == 2 and isinstance(v, (int, float)) and \
                            r > 1 and 'count' not in str(vals[0]).lower() \
                            and '$' in str(vals[0]):
                        cell.number_format = MONEY
                if vals and isinstance(vals[0], str) and vals[0].isupper():
                    ws.cell(row=r, column=1).font = BFONT
            ws['A1'].font = Font(name='Arial', size=13, bold=True)
            for c, w in enumerate(widths or [], start=1):
                ws.column_dimensions[get_column_letter(c)].width = w
            continue
        table_sheet(ws, headers, rows, widths, money_cols=money_cols,
                    center=False)
    wb.save(path)
