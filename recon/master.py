"""Read and write the Core Returns master workbook.

Layout (matches Core_Returns_RECONCILED.xlsx from the Cowork project):
  Sheet "Core Returns": Credit Memo # | Shipment ID | Control Ticket |
      Core Part # | Claim # | Credit Amount | Age (days) | Credit Requested
      | As of: | =TODAY()
  Sheet "Unpaid Cores": same columns, only rows with no Credit Memo #.
Rows with no Shipment ID come first (in their existing order), then rows
sorted by Shipment ID, oldest first.
"""
import datetime
import re

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle

from .common import clean_key

FIELDS = ['memo', 'shipment', 'ticket', 'part', 'claim', 'amount',
          'requested']
HEADERS = ['Credit Memo #', 'Shipment ID', 'Control Ticket', 'Core Part #',
           'Claim #', 'Credit Amount', 'Age (days)', 'Credit Requested']
_HEADER_KEYS = {
    'creditmemo': 'memo', 'creditmemo#': 'memo', 'memo': 'memo',
    'shipmentid': 'shipment', 'shipment': 'shipment',
    'gcrsshipperdate': 'shipment',
    'controlticket': 'ticket', 'ticket': 'ticket',
    'coreparts#': 'part', 'corepart#': 'part', 'corepart': 'part',
    'part#': 'part', 'partnumber': 'part',
    'claim#': 'claim', 'claim': 'claim', 'warrantyclaim': 'claim',
    'creditamount': 'amount', 'amount': 'amount',
    'creditrequested': 'requested',
}
WIDTHS = {'A': 14, 'B': 22, 'C': 14, 'D': 14, 'E': 10, 'F': 14, 'G': 11,
          'H': 18, 'I': 8, 'J': 12}
HEADER_FONT = Font(name='Cambria', size=11, bold=True, color='FFFFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='FF1F4E78')
DATA_FONT = Font(name='Arial', size=10)
CENTER = Alignment(horizontal='center', vertical='center')
OVERDUE_DXF = DifferentialStyle(
    font=Font(name='Arial', size=10, bold=True, color='FF9C0006'),
    fill=PatternFill('solid', fgColor='FFF8B7B7', bgColor='FFF8B7B7'))


class MasterError(Exception):
    pass


def _hkey(value):
    return re.sub(r'[\s_.:()-]', '', str(value or '')).lower()


def _clean_shipment(v):
    if v is None or v == '':
        return None
    if isinstance(v, datetime.datetime):
        if v.hour or v.minute or v.second:
            return v.strftime('%Y-%m-%d %H:%M.%S')
        return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def _clean_requested(v):
    if v is None or v == '':
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    return str(v).strip() or None


def _clean_text(v):
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip()
    return s or None


def read_master(path):
    """Return (rows, info). rows keep the workbook order."""
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
    except Exception as e:
        raise MasterError(f'Could not open the master workbook: {e}')
    ws = None
    if 'Core Returns' in wb.sheetnames:
        ws = wb['Core Returns']
    else:
        for cand in wb.worksheets:
            first = next(cand.iter_rows(min_row=1, max_row=1,
                                        values_only=True), ())
            if any(_hkey(c) == 'controlticket' for c in first):
                ws = cand
                break
    if ws is None:
        raise MasterError(
            'The master workbook needs a "Core Returns" sheet (or a sheet '
            'with a "Control Ticket" column in row 1).')
    it = ws.iter_rows(values_only=True)
    header = next(it, ())
    cols = {}
    for i, h in enumerate(header):
        k = _HEADER_KEYS.get(_hkey(h))
        if k and k not in cols:
            cols[k] = i
    if 'ticket' not in cols:
        raise MasterError('No "Control Ticket" column found in row 1.')
    rows = []
    for vals in it:
        if not vals:
            continue

        def get(k):
            i = cols.get(k)
            return vals[i] if i is not None and i < len(vals) else None

        ticket = _clean_text(get('ticket'))
        if not ticket:
            continue
        amount = get('amount')
        if isinstance(amount, str):
            try:
                amount = float(amount.replace('$', '').replace(',', ''))
            except ValueError:
                amount = None
        claim = get('claim')
        if isinstance(claim, float) and claim == int(claim):
            claim = int(claim)
        if isinstance(claim, str):
            claim = claim.strip() or None
        rows.append({
            'memo': _clean_text(get('memo')),
            'shipment': _clean_shipment(get('shipment')),
            'ticket': ticket,
            'part': _clean_text(get('part')),
            'claim': claim,
            'amount': amount,
            'requested': _clean_requested(get('requested')),
        })
    wb.close()
    return rows, {'sheet': ws.title, 'columns': sorted(cols)}


def sort_rows(rows):
    """Blank Shipment ID first (existing order), then oldest shipment first."""
    indexed = list(enumerate(rows))
    indexed.sort(key=lambda p: (0, '', p[0]) if not p[1].get('shipment')
                 else (1, str(p[1]['shipment']), p[0]))
    return [r for _, r in indexed]


def _write_sheet(ws, rows):
    for col, title in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
    ws['I1'] = 'As of:'
    ws['I1'].font = Font(name='Cambria', size=11, bold=True)
    ws['J1'] = '=TODAY()'
    ws['J1'].number_format = 'yyyy\\-mm\\-dd'
    ws.row_dimensions[1].height = 21.75
    for r, row in enumerate(rows, start=2):
        vals = [row.get('memo'), row.get('shipment'), row.get('ticket'),
                row.get('part'), row.get('claim'), row.get('amount'),
                f'=IF(B{r}="","",IFERROR($J$1-DATEVALUE(LEFT(B{r},10)),""))',
                row.get('requested')]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = DATA_FONT
        ws.cell(row=r, column=6).number_format = '"$"#,##0.00'
        ws.cell(row=r, column=7).number_format = '0'
    last = max(2, len(rows) + 1)
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:H{last}'
    ws.conditional_formatting.add(
        f'A2:G{last}',
        FormulaRule(formula=['AND($A2="",ISNUMBER($G2),$G2>60)'],
                    font=OVERDUE_DXF.font, fill=OVERDUE_DXF.fill))


def write_master(rows, path):
    rows = sort_rows(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Core Returns'
    _write_sheet(ws, rows)
    _write_sheet(wb.create_sheet('Unpaid Cores'),
                 [r for r in rows if not r.get('memo')])
    wb.save(path)
    return rows


def ticket_key(ticket):
    return clean_key(ticket)
