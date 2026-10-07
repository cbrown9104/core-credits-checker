"""Read DealerCONNECT Global Core Returns payment data.

Source screen: DealerCONNECT > Parts > Returns > Global Core Returns >
Core Tracking/Inquiry > Payment Inquiry. Columns: Tracking Number, Part
Number, Warranty Claim, Amount, Received Date, Payment Type, Document.
Accepts CSV, XLSX, or rows pasted straight from the screen (tab separated).
A row with a Document (or Received Date / Payment Type) is PAID; a row
without one is UNPAID.
"""
import csv
import datetime
import io
import re

from openpyxl import load_workbook

from .common import claim_out, clean_key, norm_part_text, parse_money

COLUMNS = ['ticket', 'part', 'claim', 'amount', 'received', 'ptype',
           'document']
_HEAD = {
    'trackingnumber': 'ticket', 'tracking': 'ticket',
    'controlticket': 'ticket', 'ticket': 'ticket',
    'partnumber': 'part', 'part': 'part', 'corepart': 'part',
    'warrantyclaim': 'claim', 'claim': 'claim', 'claimnumber': 'claim',
    'amount': 'amount', 'coreamount': 'amount',
    'receiveddate': 'received', 'received': 'received',
    'paymenttype': 'ptype',
    'document': 'document', 'documentnumber': 'document',
}


class DCError(Exception):
    pass


def _hk(v):
    return re.sub(r'[^a-z]', '', str(v or '').lower())


def _row(vals, cols):
    def get(k):
        i = cols.get(k)
        return vals[i] if i is not None and i < len(vals) else None

    ticket = clean_key(get('ticket'))
    if not ticket or ticket in ('TRACKING NUMBER', 'TOTAL'):
        return None
    received = get('received')
    if isinstance(received, (datetime.datetime, datetime.date)):
        received = received.strftime('%Y-%m-%d')
    received = str(received).strip() if received not in (None, '') else ''
    document = str(get('document') or '').strip()
    ptype = str(get('ptype') or '').strip()
    return {
        'ticket': ticket,
        'part': norm_part_text(get('part')),
        'claim': claim_out(get('claim')),
        'amount': parse_money(get('amount')),
        'received': received,
        'ptype': ptype,
        'document': document,
        'status': 'PAID' if (document or received or ptype) else 'UNPAID',
    }


def _from_table(table, source):
    """table: list of row lists. Finds the header row in the first 15."""
    cols = None
    start = 0
    for i, vals in enumerate(table[:15]):
        keys = {_HEAD.get(_hk(v)) for v in vals}
        if 'ticket' in keys and ('amount' in keys or 'part' in keys):
            cols = {}
            for j, v in enumerate(vals):
                k = _HEAD.get(_hk(v))
                if k and k not in cols:
                    cols[k] = j
            start = i + 1
            break
    if cols is None:
        # Pasted rows without a header: assume the screen's column order
        first = next((r for r in table if any(str(c).strip() for c in r)),
                     None)
        if first and re.match(r'^(C[ABW]?\d{7,9}|RA-\d+)$',
                              clean_key(first[0])):
            cols = {k: i for i, k in enumerate(COLUMNS)}
        else:
            raise DCError(
                f'{source}: could not find the DealerCONNECT columns '
                '(Tracking Number, Part Number, Warranty Claim, Amount, '
                'Received Date, Payment Type, Document).')
    out = []
    for vals in table[start:]:
        r = _row(list(vals), cols)
        if r:
            out.append(r)
    return out


def read_dc_file(path, source=''):
    low = path.lower()
    if low.endswith(('.xlsx', '.xlsm')):
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = []
        errors = []
        for ws in wb.worksheets:
            table = [list(r) for r in ws.iter_rows(values_only=True)]
            try:
                rows.extend(_from_table(table, f'{source} / {ws.title}'))
            except DCError as e:
                errors.append(str(e))
        wb.close()
        if not rows and errors:
            raise DCError(errors[0])
        return rows
    if low.endswith('.xls'):
        raise DCError(f'{source}: old .xls format — open it in Excel and '
                      'save as .xlsx or .csv, then upload again.')
    with open(path, 'r', encoding='utf-8-sig', errors='replace') as fh:
        return read_dc_text(fh.read(), source)


def read_dc_text(text, source='pasted rows'):
    text = (text or '').strip('\n')
    if not text.strip():
        return []
    sample = text[:2000]
    delim = '\t' if sample.count('\t') >= sample.count(',') else ','
    table = list(csv.reader(io.StringIO(text), delimiter=delim))
    return _from_table(table, source)


def merge(rows):
    """One record per ticket; a PAID row wins over an UNPAID one."""
    out = {}
    for r in rows:
        cur = out.get(r['ticket'])
        if cur is None or (cur['status'] == 'UNPAID' and
                           r['status'] == 'PAID'):
            out[r['ticket']] = r
    return out
